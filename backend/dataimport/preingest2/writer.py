"""
Workbook writer for Stage 7 output — the fixed Fund Master layout (TFAI shape).

One tab per schema sheet in fixed order (title banner + header + rows), plus
trailing _Audit and _Reconciliation tabs. Values are written as produced by the
assembler (copied cells, code aggregations, or formula.html derivations).
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import schema as sch

_TITLE = Font(size=14, bold=True, color='1F3A5F')
_HFILL = PatternFill('solid', fgColor='1F3A5F')
_HFONT = Font(color='FFFFFF', bold=True)


def _clean(v):
    if type(v).__name__ == 'Decimal':
        return float(v)
    return v


def _write(ws, title, columns, rows):
    ws['A1'] = title
    ws['A1'].font = _TITLE
    for ci, h in enumerate(columns, start=1):
        c = ws.cell(row=3, column=ci, value=h)
        c.fill = _HFILL
        c.font = _HFONT
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for ri, row in enumerate(rows, start=4):
        for ci, val in enumerate(row, start=1):
            ws.cell(row=ri, column=ci, value=_clean(val))
    for ci, h in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = min(max(12, len(str(h)) + 2), 34)
    if columns:
        ws.freeze_panes = ws.cell(row=4, column=1)


def build_workbook(assembled, audit, reconciliation):
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in sch.SHEETS:
        b = assembled.get(sheet['key'], {})
        ws = wb.create_sheet(sheet['name'][:31])
        _write(ws, b.get('title', sheet['title']), b.get('columns', []), b.get('rows', []))

    aws = wb.create_sheet('_Audit')
    _write(aws, 'Consolidation Audit — per file',
           ['file', 'file_class', 'entity', 'cache', 'status', 'notes'],
           [[a.get('file'), a.get('file_class'), a.get('entity'), a.get('cache'),
             a.get('status'), a.get('notes')] for a in (audit or [])])

    rws = wb.create_sheet('_Reconciliation')
    _write(rws, 'Reconciliation — tie-outs (accuracy gate)',
           ['check', 'status', 'detail'],
           [[r.get('check'), r.get('status'), r.get('detail')] for r in (reconciliation or [])])
    return wb
