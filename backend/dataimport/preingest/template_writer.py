"""
Write the assembled template to a Fund Master Workbook (TFAI.xlsx shape).

One tab per template sheet, in the fixed order, each with a title banner, a
header row, and the data rows. A trailing `_Audit` tab records coverage and the
review items so nothing is hidden — but the primary tabs are the clean fund
sheets, exactly like TFAI.xlsx.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import output_template as tmpl

_TITLE_FONT = Font(size=14, bold=True, color='1F3A5F')
_HDR_FILL = PatternFill('solid', fgColor='1F3A5F')
_HDR_FONT = Font(color='FFFFFF', bold=True)
_AUDIT_FILL = PatternFill('solid', fgColor='FDE9D9')


def _clean(v):
    if type(v).__name__ == 'Decimal':
        return float(v)
    return v


def _write_sheet(ws, title, columns, rows):
    ws['A1'] = title
    ws['A1'].font = _TITLE_FONT
    hdr_r = 3
    for ci, h in enumerate(columns, start=1):
        c = ws.cell(row=hdr_r, column=ci, value=h)
        c.fill = _HDR_FILL
        c.font = _HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for ri, row in enumerate(rows, start=hdr_r + 1):
        for ci, val in enumerate(row, start=1):
            ws.cell(row=ri, column=ci, value=_clean(val))
    for ci, h in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = min(max(12, len(str(h)) + 2), 34)
    if columns:
        ws.freeze_panes = ws.cell(row=hdr_r + 1, column=1)


def build_workbook(assembled, audit):
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in tmpl.SHEETS:
        built = assembled.get(sheet['key'], {})
        ws = wb.create_sheet(sheet['name'][:31])
        _write_sheet(ws, built.get('title', sheet['title']),
                     built.get('columns', []), built.get('rows', []))
    # trailing audit tab
    aws = wb.create_sheet('_Audit')
    aws['A1'] = 'Consolidation Audit — coverage & review'
    aws['A1'].font = _TITLE_FONT
    cols = ['file', 'kind', 'sheets_total', 'records_extracted', 'status', 'notes']
    for ci, h in enumerate(cols, start=1):
        c = aws.cell(row=3, column=ci, value=h)
        c.fill = _HDR_FILL
        c.font = _HDR_FONT
    for ri, a in enumerate(audit or [], start=4):
        for ci, h in enumerate(cols, start=1):
            aws.cell(row=ri, column=ci, value=a.get(h))
    return wb
