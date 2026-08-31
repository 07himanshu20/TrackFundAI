"""Phase D — fund LPA economic TERMS extractor (management fee, carry, hurdle,
catch-up, waterfall, clawback).

The one reframing that governs this module (v2 build rule #8 — no bare numbers
downstream): a term is NOT a scalar. "2%" alone is meaningless and dangerous — a
2% fee on committed capital vs on NAV differs several-fold in rupees. So the unit
extracted is a BUNDLE:

    FundTerm{ concept, value, unit, base, phase, text_value, source cell, verdict }

Two disciplines follow, both fail-closed:
  • UNIT is read from the cell's VALUE and its NUMBER FORMAT together. A %-formatted
    cell stores 0.02 for 2% and 1 for 100% — the format is the disambiguator (a bare
    "1" is 100% only because the cell is %-formatted). A bare number in a plain cell
    whose unit cannot be confirmed is HELD, never guessed.
  • QUALIFIERS (base, phase) are captured when stated and marked UNSPECIFIED when
    genuinely absent — never defaulted. Dropping "of committed capital, during the
    investment period" silently manufactures a wrong figure.

The model is OFF here (build rules #1-3): terms are found by a deterministic
label-anchored scan. A model fallback for free-PROSE term sheets is deferred until
we measure misses on real prose-style sheets (the one real terms sheet is a clean
key-value list — no miss to feed a model yet).

Verification (build rule #10 — every check materialises): unit-normalisation,
per-concept sanity BANDS (a value outside its band is a mis-bind → held), multi-source
AGREEMENT (same term in two places must agree, else hold-and-disclose), and
fee-vs-actual reconciliation where an actual fee figure exists.

Deferred (deliberately, proportionate scope): stepped fee-schedule computation,
per-class fee tables, and the actual waterfall/carry math — those come after NAV and
consume the scalars banked here.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from decimal import Decimal
from typing import List, Optional, Tuple

from . import lexicon, reconcile
from .cir import Provenance, Record
from .contract import TOLERANCES
from .quantity import format_pct, to_decimal

_FEE_TOL = Decimal('0.02')          # fee-vs-actual: 2% fractional tolerance (approx figures)

# ── the term catalogue: controlled synonyms + sanity bands (NEVER auto-learned;
#    a new synonym is a human-confirmed review decision, build rule #14) ────────
_UNSPEC = 'UNSPECIFIED'


@dataclass
class _Spec:
    concept: str
    kind: str                       # 'rate' | 'label'
    labels: Tuple[str, ...]         # normalised synonyms (word-boundary match)
    band: Optional[Tuple[Decimal, Decimal]] = None   # (lo, hi) fraction band for a rate
    # Which qualifiers APPLY to this concept — the n/a vs UNSPECIFIED distinction. An
    # applicable qualifier that is absent resolves to UNSPECIFIED (a disclosed, hold-worthy
    # HOLE); a genuinely-inapplicable one is 'n/a'. So a future fund whose carry DOES carry
    # a gain base flags UNSPECIFIED rather than being silently stamped n/a. `base_kind`
    # names WHICH base applies, so the committed/invested detector runs only where it fits.
    applies_base: bool = False
    applies_phase: bool = False
    base_kind: str = ''             # 'fee_base' (committed/invested/nav) | 'gain_base' | ''


# Each rate concept lists BOTH the bare label and its '% '-suffixed form. Since normalise_label now
# preserves '%' as a 'pct' token (it is the signal, see lexicon), a header like 'Carry %' normalises to
# 'carry pct'; without the '%'-form synonym its label_coverage is diluted by the extra token and a longer
# heading that merely mentions the concept ('Performance fee (carry) terms:') can out-rank the actual
# rate row. The '%'-forms restore full coverage — the same authoring ownership_pct already uses
# ('ownership %'/'stake %'). Only the RATE concepts get them (a '%' on a label term is meaningless).
_CATALOGUE = [
    _Spec('management_fee', 'rate',
          ('management fee', 'mgmt fee', 'amc fee', 'asset management fee', 'investment management fee',
           'management fee %', 'mgmt fee %', 'amc fee %'),
          band=(Decimal('0.0025'), Decimal('0.035')),
          applies_base=True, applies_phase=True, base_kind='fee_base'),
    _Spec('carried_interest', 'rate',
          ('carried interest', 'carry', 'performance fee', 'perf fee', 'incentive fee', 'profit share',
           'carried interest %', 'carry %', 'performance fee %', 'incentive fee %'),
          band=(Decimal('0.05'), Decimal('0.30')),
          applies_base=True, base_kind='gain_base'),   # carry base = total vs realized gains (no detector yet → UNSPECIFIED)
    _Spec('hurdle_rate', 'rate',
          ('hurdle', 'preferred return', 'pref return', 'hurdle rate', 'irr hurdle',
           'hurdle %', 'hurdle rate %', 'preferred return %', 'pref return %'),
          band=(Decimal('0.04'), Decimal('0.12'))),
    _Spec('catch_up', 'rate',
          ('catch up', 'catch-up', 'gp catch up', 'catchup', 'catch up %', 'catch-up %', 'catchup %'),
          band=(Decimal('0'), Decimal('1'))),
    _Spec('clawback_holdback', 'rate',
          ('clawback', 'holdback', 'clawback holdback reserve', 'holdback reserve',
           'clawback %', 'holdback %'),
          band=(Decimal('0'), Decimal('0.50'))),
    _Spec('waterfall', 'label',
          ('waterfall', 'distribution waterfall')),
]
_SPEC = {s.concept: s for s in _CATALOGUE}

# base/phase keyword maps (matched only in a fee CONTEXT, so a stray 'committed' elsewhere
# never binds a base). value → canonical token.
_BASE_KEYWORDS = (('committed capital', 'committed_capital'), ('committed', 'committed_capital'),
                  ('invested capital', 'invested_capital'), ('invested', 'invested_capital'),
                  ('drawn capital', 'drawn_capital'), ('drawn', 'drawn_capital'),
                  ('net asset value', 'nav'), ('nav', 'nav'))
_PHASE_KEYWORDS = (('investment period', 'investment_period'), ('commitment period', 'investment_period'),
                   ('post investment', 'post_investment'), ('post-investment', 'post_investment'),
                   ('after the investment period', 'post_investment'))


@dataclass
class FundTerm:
    """One LPA economic term as a self-describing bundle (build rule #8)."""
    concept: str
    value: Optional[Decimal]        # RATE as a FRACTION (0.02 == 2%); None for a label term
    unit: str                       # 'percent' | 'label' | 'UNSPECIFIED'
    base: str                       # committed_capital | invested_capital | nav | n/a | UNSPECIFIED
    phase: str                      # investment_period | post_investment | whole_life | n/a | UNSPECIFIED
    text_value: Optional[str]       # e.g. 'European whole fund' for a label term
    provenance: Provenance
    verdict: str = 'confirmed'      # confirmed | held | gap
    hold_reason: str = ''
    qualifiers: dict = field(default_factory=dict)

    @property
    def confirmed(self) -> bool:
        return self.verdict == 'confirmed'

    @property
    def display(self) -> str:
        if self.text_value is not None:
            return self.text_value
        if self.value is None:
            return '—'
        return _fmt_pct(self.value)


# ── small helpers ──────────────────────────────────────────────────────────
def _a1(col: Optional[int], row: Optional[int]) -> str:
    if col is None or row is None:
        return ''
    from openpyxl.utils import get_column_letter
    return f'{get_column_letter(col + 1)}{row + 1}'


def _norm(s: str) -> str:
    return lexicon.normalise_label(s)


def _fmt_pct(frac: Decimal) -> str:
    return format_pct(frac)          # canonical; see quantity.format_pct (was a 3rd hand-rolled copy)


def _has_syn(label, synonyms) -> bool:
    """Whole-word, normalise-both-sides match — delegates to the ONE shared matcher so this
    module never hand-rolls its own (that is how the raw-vs-normalised/substring bug classes
    recur). Takes the RAW label; the matcher normalises internally."""
    return lexicon.label_matches_any(label, synonyms)


def _num(s: str) -> Optional[Decimal]:
    m = re.search(r'-?\d[\d,]*\.?\d*', s.replace('–', '-'))
    return to_decimal(m.group(0).replace(',', '')) if m else None


def _value_format_grid(path: str, sheet_name: str):
    """A (value, number_format) grid for ONE sheet — read fresh so value and format
    are aligned (the profiler grid is data_only and drops formats)."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            return []
        return [[(c.value, c.number_format) for c in row] for row in wb[sheet_name].iter_rows()]
    finally:
        wb.close()


# ── unit normalisation — VALUE + NUMBER FORMAT together (the core rule) ──────
def normalize_percent(raw, number_format: str, raw_label: str) -> Tuple[Optional[Decimal], str, str]:
    """Return (fraction, unit, note). fraction is the rate as a fraction (2% → 0.02).
    Fail-closed: a bare number whose unit cannot be confirmed → (None, UNSPECIFIED, reason)."""
    fmt = number_format or ''
    label_has_pct = ('%' in (raw_label or '')) or ('percent' in _norm(raw_label or ''))
    if isinstance(raw, str):
        s = raw.strip().lower()
        if 'bps' in s or 'basis point' in s:
            n = _num(s)
            return (n / Decimal(10000), 'percent', 'from bps') if n is not None else (None, _UNSPEC, 'bps parse failed')
        if '%' in s:
            n = _num(s)
            return (n / Decimal(100), 'percent', 'from % string') if n is not None else (None, _UNSPEC, '% parse failed')
        n = _num(s)
        if n is None:
            return None, _UNSPEC, 'non-numeric value'
        raw = n
    d = to_decimal(raw)
    if d is None:
        return None, _UNSPEC, 'non-numeric value'
    if '%' in fmt:                                   # %-formatted cell stores the fraction
        return d, 'percent', 'percent-format cell'
    if Decimal(0) < d < Decimal(1):                 # plain cell, fraction magnitude → a fraction
        return d, 'percent', 'inferred fraction (0<v<1)'
    if d >= 1 and label_has_pct:                     # plain cell, label carries % → value is the percent
        return d / Decimal(100), 'percent', 'inferred bare-number %-label'
    return None, _UNSPEC, f'bare number {d} in plain format — unit unconfirmable'


# ── locate + read one concept on a sheet ────────────────────────────────────
def _find_label(grid, synonyms):
    """The TIGHTEST label match — the cell the synonym occupies most fully — so a
    row that IS the label ('Carry %') out-anchors a section HEADING that merely
    mentions it ('Management & performance fee working'). Ties break to first seen.
    Coverage/whole-word matching comes from the ONE shared matcher (lexicon)."""
    best = None                                     # (coverage, r, c, text)
    for r, row in enumerate(grid):
        for c, (v, _f) in enumerate(row):
            if not isinstance(v, str):
                continue
            cov = lexicon.label_coverage(v, synonyms)
            if cov > 0 and (best is None or cov > best[0]):
                best = (cov, r, c, v)
    return (best[1], best[2], best[3]) if best else (None, None, None)


def _resolve_value(grid, lr, lc):
    """value cell: look right → look below → parse the label cell's own sentence."""
    for c in range(lc + 1, len(grid[lr])):
        v, f = grid[lr][c]
        if v not in (None, ''):
            return v, f, lr, c
    if lr + 1 < len(grid) and lc < len(grid[lr + 1]):
        v, f = grid[lr + 1][lc]
        if v not in (None, ''):
            return v, f, lr + 1, lc
    lv, lf = grid[lr][lc]
    if isinstance(lv, str) and ('%' in lv or _num(lv) is not None):
        return lv, lf, lr, lc
    return None, None, None, None


def _classify_waterfall(text: str) -> str:
    low = text.lower()
    if 'american' in low or 'deal by deal' in low or 'deal-by-deal' in low:
        return 'american_deal_by_deal'
    if 'european' in low or 'whole fund' in low or 'whole-fund' in low:
        return 'european_whole_fund'
    return 'other'


def _resolve_base_phase(grid, spec: _Spec):
    """Base/phase for a concept, honouring per-concept APPLICABILITY:
      • a qualifier that does NOT apply to this concept → 'n/a' (structural absence);
      • a qualifier that APPLIES but is not found → UNSPECIFIED (a disclosed hole),
        never defaulted.
    The committed/invested/nav detector runs only for a 'fee_base' concept (management
    fee); a 'gain_base' concept (carry) has no deterministic detector in this slice, so
    an applicable gain base is UNSPECIFIED — correctly flagged, not stamped n/a."""
    base = _UNSPEC if spec.applies_base else 'n/a'
    phase = _UNSPEC if spec.applies_phase else 'n/a'
    src = None
    if spec.base_kind == 'fee_base':                 # committed/invested/nav — mgmt-fee-specific
        for r, row in enumerate(grid):
            for c, (v, _f) in enumerate(row):
                if not isinstance(v, str):
                    continue
                low = v.lower()
                if not ('fee' in low or 'management' in low or '%' in v or 'base' in low):
                    continue
                if base == _UNSPEC:
                    for kw, tok in _BASE_KEYWORDS:
                        if kw in low:
                            base, src = tok, (r, c)
                            break
                if spec.applies_phase and phase == _UNSPEC:
                    for kw, tok in _PHASE_KEYWORDS:
                        if kw in low:
                            phase = tok
                            break
    return base, phase, src


def _extract_concept(spec, grid, *, source_label, sheet, content_fp) -> Optional[FundTerm]:
    lr, lc, label_text = _find_label(grid, spec.labels)
    if lr is None:
        return None                                 # concept not on this sheet
    raw, fmt, vr, vc = _resolve_value(grid, lr, lc)
    prov = Provenance(source_file=source_label, content_fingerprint=content_fp, sheet=sheet,
                      cell=_a1(vc, vr), row_label=label_text)
    if raw is None:
        return FundTerm(spec.concept, None, _UNSPEC, 'n/a', 'n/a', None, prov,
                        verdict='gap', hold_reason=f'{spec.concept}: label found, no value')

    if spec.kind == 'label':
        text = str(raw).strip()
        wf = _classify_waterfall(text)
        return FundTerm(spec.concept, None, 'label', 'n/a', 'whole_life' if wf == 'european_whole_fund' else 'n/a',
                        text, prov, verdict='confirmed', qualifiers={'waterfall_type': wf})

    frac, unit, note = normalize_percent(raw, fmt, label_text)
    base, phase, bsrc = _resolve_base_phase(grid, spec)
    quals = {'unit_note': note}
    if bsrc is not None:
        quals['base_source'] = _a1(bsrc[1], bsrc[0])
    if spec.concept == 'catch_up':
        quals['catch_up_exists'] = frac is not None and frac > 0
    if spec.concept == 'hurdle_rate':
        # compounding (simple/compounded) is an APPLICABLE qualifier that lives in LPA
        # PROSE ("compounded annually"), not a structured cell — so deterministically it
        # is UNSPECIFIED (a disclosed hole), the cleanest first case for the model
        # prose-fallback. Never defaulted to 'simple'.
        quals['compounding'] = _UNSPEC
    if frac is None:                                # unit unconfirmable → HELD (fail-closed)
        return FundTerm(spec.concept, None, _UNSPEC, base, phase, None, prov,
                        verdict='held', hold_reason=f'{spec.concept}: {note}', qualifiers=quals)
    if spec.band is not None and not (spec.band[0] <= frac <= spec.band[1]):
        return FundTerm(spec.concept, frac, unit, base, phase, None, prov, verdict='held',
                        hold_reason=f'{spec.concept}: {_fmt_pct(frac)} outside band '
                                    f'[{_fmt_pct(spec.band[0])},{_fmt_pct(spec.band[1])}] — likely mis-bound',
                        qualifiers=quals)
    return FundTerm(spec.concept, frac, unit, base, phase, None, prov, verdict='confirmed', qualifiers=quals)


def _terms_sheet(prof) -> Optional[str]:
    """The sheet bearing the most economic-term labels (≥2 distinct concepts)."""
    best, best_n = None, 0
    for s in prof['sheets']:
        rows = prof['grid'][s.sheet]
        found = set()
        for row in rows:
            for v in row:
                if isinstance(v, str):
                    for spec in _CATALOGUE:
                        if _has_syn(v, spec.labels):
                            found.add(spec.concept)
        if len(found) > best_n:
            best, best_n = s.sheet, len(found)
    return best if best_n >= 2 else None


# fund-identity labels (name, legal structure/SEBI category, vintage) — plain descriptive
# fields, read label-anchored from the terms sheet via the SHARED lexicon matcher (model OFF).
_ID_FUND_NAME = ('fund name', 'scheme name', 'name of the fund', 'name of fund', 'name of the scheme')
_ID_LEGAL = ('legal structure', 'legal form', 'fund structure', 'constitution', 'aif category', 'category')
_ID_VINTAGE = ('vintage', 'vintage year', 'vintage yr', 'inception year', 'year of inception', 'first close year')


def _read_fund_identity(grid) -> dict:
    """Fund name / legal structure (SEBI category) / vintage — the value in the cell to the
    RIGHT of a matching label. Whole-word match via the one shared matcher, never a copy."""
    out: dict = {}
    for row in grid:
        vals = [c[0] if isinstance(c, tuple) else c for c in row]     # cells are (value, number_format)
        lab_idx = next((i for i, v in enumerate(vals) if isinstance(v, str) and str(v).strip()), None)
        if lab_idx is None:
            continue
        label = vals[lab_idx]
        val = next((v for v in vals[lab_idx + 1:] if v not in (None, '') and str(v).strip()), None)
        if val is None:
            continue
        if 'fund_name' not in out and lexicon.label_matches_any(label, _ID_FUND_NAME):
            out['fund_name'] = str(val).strip()
        elif 'legal_structure' not in out and lexicon.label_matches_any(label, _ID_LEGAL):
            out['legal_structure'] = str(val).strip()
        elif 'vintage_year' not in out and lexicon.label_matches_any(label, _ID_VINTAGE):
            out['vintage_year'] = str(val).strip()
    return out


def extract_fund_terms(label, path, prof, *, content_fp='') -> Optional[Record]:
    """Extract the economic-term BUNDLES from ONE fund file's terms-bearing sheet.
    Returns a Record(domain='fund_terms', entity_id='fund') or None if the file has
    no terms sheet. Each field is a FundTerm bundle (build rule #8); fund-identity fields
    (name/legal-structure/vintage) are attached as plain strings alongside."""
    sheet = _terms_sheet(prof)
    if sheet is None:
        return None
    grid = _value_format_grid(path, sheet)
    if not grid:
        return None
    fields: dict = {'fund': 'fund'}
    for spec in _CATALOGUE:
        term = _extract_concept(spec, grid, source_label=label, sheet=sheet, content_fp=content_fp)
        if term is not None:
            fields[spec.concept] = term
    fields.update(_read_fund_identity(grid))         # descriptive identity (plain strings)
    if len(fields) == 1:                             # only the 'fund' scalar → nothing found
        return None
    return Record('fund_terms', entity_id='fund', fields=fields)


# ── cross-file reconcile: multi-source agreement + base/phase enrichment ─────
_RATE_TOL = Decimal('0.0005')       # rates must match to 0.05 percentage-point

_FEE_LABEL_SYN = ('management fee', 'mgmt fee', 'amc fee', 'asset management fee')
_PERIOD_BALANCE = ('payable', 'accrued', 'provision', 'balance', 'outstanding')
_PERIOD_ITD = ('itd', 'inception', 'cumulative', 'to date', 'to-date', 'life to date')
_PERIOD_ANNUAL = ('annual', 'per annum', 'p.a.', 'per year')


def _fee_period(text: str) -> str:
    """Classify a fee figure's PERIOD from its label — the period-select disambiguation
    (annual vs ITD vs balance) that fee-vs-actual needs and NAV will reuse."""
    low = text.lower()
    if any(k in low for k in _PERIOD_BALANCE):
        return 'balance'
    if any(k in low for k in _PERIOD_ITD):
        return 'itd'
    if any(k in low for k in _PERIOD_ANNUAL):
        return 'annual'
    return 'unknown'


def resolve_actual_annual_fee(label, prof, *, content_fp='') -> Tuple[Optional[Decimal], Optional[Provenance], List[dict]]:
    """The ANNUAL ACTUAL management fee, selected INDEPENDENTLY of any expected value
    (so a wrong actual reddens the tie-out). Prefers a budget-vs-ACTUAL sheet's 'actual'
    column on the mgmt-fee row — an annual figure independent of the rate×base working —
    and EXCLUDES ITD/cumulative and payable/balance figures by their period label.
    Returns (value, provenance, all_candidates) where all_candidates discloses every fee
    figure seen with its period classification (so the selection is auditable)."""
    candidates: List[dict] = []
    chosen = (None, None)
    for s in prof['sheets']:
        rows = prof['grid'][s.sheet]
        # locate a budget-vs-actual header (an 'actual' column, ideally beside 'budget')
        actual_col = hdr = None
        for r, row in enumerate(rows):
            cols = {_norm(v): c for c, v in enumerate(row) if isinstance(v, str)}
            if 'actual' in cols and 'budget' in cols:
                actual_col, hdr = cols['actual'], r
                break
        if actual_col is None:
            continue
        for r in range(hdr + 1, len(rows)):
            row = rows[r]
            lbl = next((v for v in row if isinstance(v, str)), '')
            if not _has_syn(lbl, _FEE_LABEL_SYN):
                continue
            period = _fee_period(lbl)
            val = to_decimal(row[actual_col]) if actual_col < len(row) else None
            if val is None:
                continue
            prov = Provenance(source_file=label, content_fingerprint=content_fp, sheet=s.sheet,
                              cell=_a1(actual_col, r), row_label=lbl)
            candidates.append({'value': str(val), 'period': period, 'cell': f'{s.sheet}!{prov.cell}'})
            if period not in ('balance', 'itd') and chosen[0] is None:   # annual/unknown-annual actual
                chosen = (val, prov)
    return chosen[0], chosen[1], candidates


# ── Reference-only fund AGGREGATE comparators (U7 soft-check RHS) ─────────────────
# A fund file may state its OWN top-down aggregate portfolio figures — total portfolio
# revenue / EBITDA, budget AND actual — typically on a budget-vs-actual or forecast sheet.
# These are the RIGHT-HAND SIDE of the portfolio-vs-fund soft correspondence: they let the
# bottom-up Σ of the company MIS files be checked against the fund's own recorded aggregate.
#
# They are REFERENCE-ONLY and NEVER emitted as authoritative actuals — that would breach the
# Budget-vs-Actual guard (a budget/forecast number must never be read as an actual). They land
# in cir.comparators (a Record-free store the assembler never reads), never in a Figure. This
# reader is the deliberate EXTENSION of that guard, not a breach: it reaches into the fenced
# sheet ONLY to draw reference comparators, tagged reference_only.
#
# Keyed on CONCEPT (revenue / EBITDA lexicon) + an AGGREGATE-ROLE marker (portfolio /
# aggregate / consolidated / total) — never a sheet NAME — so it generalises to any fund's
# file. A concept with no such line is simply absent (the reconciliation stage discloses it,
# never fabricates a comparison).
_AGG_REV_SYN = ('revenue', 'turnover', 'sales', 'total income', 'top line', 'topline',
                'income from operations', 'operating revenue', 'revenue from operations',
                'revenues from operations')
_AGG_EBITDA_SYN = ('ebitda', 'operating profit', 'operating income')
_AGG_ROLE_MARKERS = ('portfolio', 'aggregate', 'consolidated', 'investee', 'combined', 'total')
# The stated period basis lives in a PARENTHETICAL qualifier ("(aggregate, ann.)"), which
# normalise_label strips — so basis is read off the RAW label, whole-word, never the normalised form.
_ANNUALISED_RE = re.compile(r'\bann|\bp\.?\s?a\b|per\s+annum', re.I)
_REFERENCE_COMPARATORS = (
    ('portfolio_revenue', _AGG_REV_SYN),
    ('portfolio_ebitda', _AGG_EBITDA_SYN),
)


def _stated_basis(raw_label) -> str:
    """The period basis a fund aggregate line states about ITSELF, read from the raw label's
    parenthetical qualifier ('… ann.' / 'p.a.' / 'per annum'). 'annualised' or 'unspecified'
    (never guessed — an unlabelled aggregate is disclosed as basis-unspecified, not assumed annual)."""
    return 'annualised' if raw_label and _ANNUALISED_RE.search(str(raw_label)) else 'unspecified'


def extract_reference_comparators(label, prof, *, content_fp='') -> List[dict]:
    """Read a fund's OWN aggregate portfolio revenue/EBITDA (budget AND actual) as REFERENCE-ONLY
    comparators — the RHS of the U7 portfolio-vs-fund soft check. Returns a list of dicts; each is
    self-describing (concept, budget/actual value + cell, stated basis, provenance). Never mutates
    the CIR and never produces a Figure — the caller files these under cir.comparators. Empty list
    when the file states no such aggregate (absence is disclosed downstream, never fabricated)."""
    out: List[dict] = []
    for s in prof['sheets']:
        rows = prof['grid'][s.sheet]
        # a budget-vs-actual grid: a header row carrying BOTH a 'budget' and an 'actual' column.
        # (Same detection the fee-actual reader uses — the ONE idiom, not a second hand-rolled scan.)
        budget_col = actual_col = hdr = None
        for r, row in enumerate(rows):
            cols = {_norm(v): c for c, v in enumerate(row) if isinstance(v, str)}
            if 'actual' in cols and 'budget' in cols:
                budget_col, actual_col, hdr = cols['budget'], cols['actual'], r
                break
        if hdr is None:
            continue                                   # no budget-vs-actual grid here — nothing to draw
        for concept, syns in _REFERENCE_COMPARATORS:
            for r in range(hdr + 1, len(rows)):
                row = rows[r]
                lbl = next((v for v in row if isinstance(v, str)), '')
                if not (_has_syn(lbl, syns) and _has_syn(lbl, _AGG_ROLE_MARKERS)):
                    continue                           # concept + aggregate-role marker BOTH required
                bud = to_decimal(row[budget_col]) if budget_col < len(row) else None
                act = to_decimal(row[actual_col]) if actual_col < len(row) else None
                if bud is None and act is None:
                    continue
                out.append({
                    'concept': concept, 'reference_only': True,
                    'source_file': label, 'sheet': s.sheet, 'row_label': lbl,
                    'content_fp': content_fp,
                    'budget': (str(bud) if bud is not None else None),
                    'actual': (str(act) if act is not None else None),
                    'budget_cell': (f'{s.sheet}!{_a1(budget_col, r)}' if bud is not None else ''),
                    'actual_cell': (f'{s.sheet}!{_a1(actual_col, r)}' if act is not None else ''),
                    'basis': _stated_basis(lbl),
                })
                break                                   # first matching row per concept on this sheet
    return out


def extract_realised_gross(label, prof, *, content_fp='') -> Optional[dict]:
    """The fund's OWN stated realised GROSS proceeds — the concept-matched RHS for the
    realisations-vs-exits soft check (Σ of the exit ledger's gross proceeds ties to it). REFERENCE-ONLY.

    Concept-matched to GROSS specifically: realised NET (a smaller figure) and realised GAIN
    (proceeds − cost) are DIFFERENT concepts, and binding the exit-gross Σ to either would compare
    two unlike things. The 'gross' qualifier usually sits in a parenthetical ('Realised value (gross
    proceeds)') which normalise_label strips — so this matches the RAW label, and requires 'gross'
    while excluding 'net'/'gain'. Returns None when the file states no explicit realised-gross line."""
    for s in prof['sheets']:
        rows = prof['grid'][s.sheet]
        for r, row in enumerate(rows):
            lbl = next((v for v in row if isinstance(v, str)), '')
            low = str(lbl).lower()
            if not ('realis' in low or 'realiz' in low) or 'gross' not in low:
                continue                                 # realised + GROSS only (net/gain are other concepts)
            val = vcol = None
            for c, v in enumerate(row):                  # the value = first numeric cell on the row
                if isinstance(v, (int, float)):
                    val, vcol = to_decimal(v), c
                    break
            if val is None:
                continue
            return {'concept': 'realised_gross', 'reference_only': True,
                    'source_file': label, 'sheet': s.sheet, 'row_label': lbl,
                    'content_fp': content_fp, 'actual': str(val), 'budget': None,
                    'actual_cell': f'{s.sheet}!{_a1(vcol, r)}', 'budget_cell': '',
                    'basis': 'cumulative'}
    return None


def fee_base_phase(label, path, prof, *, content_fp='') -> Tuple[str, str, Optional[Provenance]]:
    """Scan a fund file for a fee's BASE/PHASE stated in a fee context (committed /
    invested / nav + investment-period). Returns (base, phase, Provenance) or
    (UNSPECIFIED, UNSPECIFIED, None). Enriches a mgmt-fee term whose own sheet did
    not state the base — the classic terms-sheet gap (rate on one sheet, base on another)."""
    for s in prof['sheets']:
        grid = [[(v, None) for v in row] for row in prof['grid'][s.sheet]]
        base, phase, src = _resolve_base_phase(grid, _SPEC['management_fee'])
        if base != _UNSPEC or phase != _UNSPEC:
            return base, phase, Provenance(source_file=label, content_fingerprint=content_fp,
                                           sheet=s.sheet, cell=_a1(src[1], src[0]) if src else '')
    return _UNSPEC, _UNSPEC, None


def reconcile_fund_terms(records, *, base_phase_hint=None, fee_actual=None,
                         committed_base=None) -> Tuple[Optional[Record], List[dict]]:
    """Merge per-file terms records into ONE canonical record and materialise checks
    (build rule #10). Disciplines:
      • MULTI-SOURCE AGREEMENT — a concept in >1 source must agree; on agreement it is
        corroborated (a second independent signal), on disagreement BOTH are HELD
        (never silently pick one — the multi-value guard);
      • BASE/PHASE ENRICHMENT — a mgmt-fee term whose base is UNSPECIFIED is filled from
        base_phase_hint (a second source that states it), with the enriching cell cited;
      • FEE-VS-ACTUAL — rate × committed_base must reproduce the ANNUAL ACTUAL fee. This
        is the check that makes a term TRUE, not merely faithfully transcribed: agreement
        proves two documents match; only this proves the stated term is what's charged. A
        PASS proves the base ARITHMETICALLY over called/invested (rate×called ≠ rate×committed
        while called < committed) within a DISCLOSED ±tolerance band; the NAV leg is closed by
        the NAV slice (rate×NAV, once NAV is extracted — early-life NAV≈committed cannot be
        separated by arithmetic alone, so 'not NAV' is not claimed here). A mismatch RETRACTS
        the label-only base to UNSPECIFIED and flags the term's verdict HELD while KEEPING the
        corroborated rate — the mismatch impugns the base, not the independently-agreed rate.
    Canonical = the source with the most confirmed terms (the dedicated terms sheet)."""
    records = [r for r in records if r is not None]
    if not records:
        return None, []
    checks: List[dict] = []

    def _nconf(rec):
        return sum(1 for v in rec.fields.values() if isinstance(v, FundTerm) and v.confirmed)
    primary = max(records, key=_nconf)
    others = [r for r in records if r is not primary]

    for concept, pt in list(primary.fields.items()):
        if not isinstance(pt, FundTerm):
            continue
        for o in others:
            ot = o.fields.get(concept)
            if not isinstance(ot, FundTerm) or ot.value is None or pt.value is None:
                continue
            cid = f'term_agrees_{concept}'
            if abs(pt.value - ot.value) <= _RATE_TOL:
                pt.qualifiers['corroborated_by'] = f'{ot.provenance.sheet}!{ot.provenance.cell}'
                checks.append(reconcile._result(cid, reconcile.HARD, reconcile.PASS,
                              lhs=_fmt_pct(pt.value), rhs=_fmt_pct(ot.value),
                              detail=f'{concept} agrees across two sources'))
            else:
                pt.verdict = ot.verdict = 'held'
                why = (f'{concept} disagrees across sources: {_fmt_pct(pt.value)} '
                       f'@{pt.provenance.sheet}!{pt.provenance.cell} vs {_fmt_pct(ot.value)} '
                       f'@{ot.provenance.sheet}!{ot.provenance.cell} — held')
                pt.hold_reason = ot.hold_reason = why[:90]
                checks.append(reconcile._result(cid, reconcile.HARD, reconcile.FAIL,
                              lhs=_fmt_pct(pt.value), rhs=_fmt_pct(ot.value), detail=why))

    if base_phase_hint is not None:
        b, ph, prov = base_phase_hint
        mt = primary.fields.get('management_fee')
        if isinstance(mt, FundTerm) and prov is not None:
            enriched = []
            if mt.base == _UNSPEC and b != _UNSPEC:
                mt.base = b
                mt.qualifiers['base_source'] = f'{prov.sheet}!{prov.cell}'
                enriched.append(f'base={b}')
            if mt.phase == _UNSPEC and ph != _UNSPEC:
                mt.phase = ph
                enriched.append(f'phase={ph}')
            if enriched:
                checks.append(reconcile._result('management_fee_base_enriched', reconcile.DISCLOSURE,
                              'disclosed', detail=f'mgmt-fee {", ".join(enriched)} from '
                              f'{prov.sheet}!{prov.cell} (terms sheet stated only the rate)'))

    # ── fee-vs-actual: the check that makes the mgmt fee TRUE, not merely transcribed ──
    mt = primary.fields.get('management_fee')
    if isinstance(mt, FundTerm) and mt.value is not None and fee_actual is not None and committed_base is not None:
        actual_val, actual_prov, _cands = fee_actual
        if actual_val is not None and actual_val != 0:
            expected = mt.value * Decimal(str(committed_base))
            band = _fmt_pct(_FEE_TOL)                     # the ± tie-out tolerance, disclosed (not implicit)
            tie = f'{_fmt_pct(mt.value)}×{committed_base}={expected} vs actual {actual_val}@{actual_prov.sheet}!{actual_prov.cell}'
            if abs(expected - actual_val) / abs(actual_val) <= _FEE_TOL:
                mt.qualifiers['fee_vs_actual'] = f'PROVEN within ±{band}: {tie}'
                mt.qualifiers['base_proven'] = 'fee_vs_actual'   # base proven by arithmetic, not just a label
                checks.append(reconcile._result('management_fee_vs_actual', reconcile.HARD, reconcile.PASS,
                              lhs=str(expected), rhs=str(actual_val),
                              detail=f'rate×committed reproduces the actual annual fee within ±{band} — base proven '
                                     f'committed over called/invested (NAV leg pending NAV slice); {tie}'))
            else:
                # arithmetic contradicts the stated terms → the BASE is not proven. Retract the
                # label-only base (arithmetic beats a single label read); flag the term's verdict
                # HELD but KEEP the corroborated rate — the mismatch impugns the base, not the
                # independently-agreed rate (proportionate retraction, not whole-term suppression).
                if mt.qualifiers.get('base_source') and 'base_proven' not in mt.qualifiers:
                    mt.base = _UNSPEC
                    mt.qualifiers['base_retracted'] = 'fee_vs_actual mismatch — label-only base unproven'
                mt.verdict = 'held'
                mt.hold_reason = (f'{_fmt_pct(mt.value)}×committed {committed_base}={expected} ≠ '
                                  f'actual {actual_val} — base mis-bound or terms overridden')[:90]
                checks.append(reconcile._result('management_fee_vs_actual', reconcile.HARD, reconcile.FAIL,
                              lhs=str(expected), rhs=str(actual_val),
                              detail=f'{tie}; deviation exceeds ±{_fmt_pct(_FEE_TOL)} tolerance — base retracted to '
                                     f'UNSPECIFIED, corroborated rate kept'))
    return primary, checks
