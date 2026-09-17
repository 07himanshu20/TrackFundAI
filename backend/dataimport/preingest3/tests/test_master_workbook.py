"""LOCK for the faithful 13-sheet master workbook (master_workbook.build_master).

Four disciplines, each with a REDDENING control (proves the test goes RED against the
exact bug it guards, per the negative-control rule — not just green against the fix):
  1. Total FV = Σ company FV → PASS on confirmed FVs.
  2. HONEST STATE — a HELD figure (even one carrying a value for the reviewer) renders
     ⚠HELD, NEVER its untrusted number. Reddens if the renderer ever shows value_cr.
  3. HONEST NAV HOLD — when NAV holds, TVPI=DPI+RVPI reads HELD, never faked to PASS.
     Reddens if the identity is forced green.
  4. DETERMINISM — same CIR → byte-identical workbook (pure function).
Plus staleness (tunable + fail-closed on an unparseable date) and one slow real-file E2E.
"""
import io
import os
from decimal import Decimal

import openpyxl
import pytest

from backend.dataimport.preingest3 import master_workbook as mw
from backend.dataimport.preingest3.cir import CIR, Record, Figure, Provenance
from backend.dataimport.preingest3.quantity import format_pct


def _pv(cell='S!A1', col=''):
    return Provenance(source_file='t', content_fingerprint='', sheet='S', cell=cell, col_label=col)


def _fig(concept, val, *, held=False, gap=False, basis='point_in_time', col=''):
    return Figure(concept, None if (gap or val is None) else Decimal(str(val)), None, _pv(col=col),
                  held=held, gap=gap, basis=basis, hold_reason='ambiguous period' if held else '')


def _inv(name, cost, fv, *, sector='Tech', fv_held=False):
    return Record('portfolio_investments', entity_id=name, fields={
        'company': name, 'sector': sector, 'stage': 'Growth', 'investment_date': '30-Sep-21',
        'ownership_pct': 0.2, 'cost': _fig('cost', cost),
        'fair_value': _fig('fair_value', fv, held=fv_held)})


def _base_cir(*, nav_held=False, as_of='2026-06-30'):
    """A minimal but complete fund CIR: 2 investments, capital account, NAV (held or not)."""
    cir = CIR(as_of=as_of, rate_card_id='rc_test')
    cir.records.append(_inv('Alpha Co', 100, 100))
    cir.records.append(_inv('Beta Co', 200, 200))
    cir.records.append(Record('fund_financials', entity_id='fund',
                              fields={'called': _fig('called', 600, basis='cumulative'),
                                      'distributed': _fig('distributed', 70, basis='cumulative')}))
    cir.records.append(Record('nav', entity_id='fund', fields={
        'lp_nav': _fig('lp_nav', 806.25, held=nav_held),
        'gross_nav': _fig('gross_nav', 867.85, held=nav_held),
        'as_of': as_of, 'derivation': 'called − distributed + net ITD P&L',
        'components': [('called', '600', 'Drawdowns!D11')],
        'carry_value': '61.6', 'carry_ceiling': '84.8',
        'hold_reason': 'as-of cutoffs disagree' if nav_held else ''}))
    return cir


def _sheet_rows(wb, name):
    return [tuple('' if c is None else c for c in r) for r in wb[name].iter_rows(values_only=True)]


def _find(rows, col0_startswith):
    for r in rows:
        if r and isinstance(r[0], str) and r[0].startswith(col0_startswith):
            return r
    return None


# 1 ── Total FV = Σ company FV → PASS ────────────────────────────────────────
def test_fig_basis_discloses_partial_year_span_REDDENING():
    # #3 Q1: a PARTIAL-year figure must be unmissable beside full-year peers. Analisa emits a TRUE
    # 5-month YTD (₹6.87Cr) — we DISCLOSE the span "YTD · 5mo (Jan–May)" rather than destroy the true
    # number or fake an annual. Reddening: the old bare "YTD" could be read as annual next to peers.
    rev = _fig('revenue', 6.87, basis='YTD'); rev.months = 5
    figs = {'revenue': rev, 'ebitda': None, 'cash': None, 'headcount': None}
    assert mw._fig_basis(figs, mo=5) == 'YTD · 5mo (Jan–May)'    # span + reporting month → month range
    assert mw._fig_basis(figs, mo=None) == 'YTD · 5mo'           # no reporting month → span count alone
    # must-not-misfire: a full-year (12mo) figure shows the bare basis, no span suffix
    fy = _fig('revenue', 50, basis='YTD'); fy.months = 12
    assert mw._fig_basis({'revenue': fy, 'ebitda': None, 'cash': None, 'headcount': None}, mo=12) == 'YTD'
    # must-not-misfire: a point-in-time stock (months 0) shows the bare basis, no span
    pit = _fig('cash', 5, basis='point_in_time'); pit.months = 0
    assert mw._fig_basis({'revenue': None, 'ebitda': None, 'cash': pit, 'headcount': None}, mo=5) \
        == 'point_in_time'


def test_total_fv_identity_passes():
    wb = mw.build_master(_base_cir(), files=['x'])
    row = _find(_sheet_rows(wb, 'MOIC_TVPI_DPI'), 'Total FV = Σ company FV')
    assert row is not None
    assert row[1] == 300.0 and row[2] == 300.0            # Σ FV = 100 + 200
    assert 'PASS' in row  # verdict column


# 2 ── HONEST STATE + reddening: a held figure never shows its untrusted value ─
def test_held_figure_renders_honest_state_not_value_REDDENING():
    cir = _base_cir()
    # a company MIS whose revenue is HELD but CARRIES a value (99) for the reviewer
    cir.records.append(Record('mis', entity_id='Alpha Co', fields={
        'company': 'Alpha Co',
        'revenue': _fig('revenue', 99, held=True, col='2026-02-28'),
        'ebitda': _fig('ebitda', 5, col='2026-02-28'),
        'cash': _fig('cash', None, gap=True), 'headcount': _fig('headcount', None, gap=True)}))
    wb = mw.build_master(cir, files=['x'])
    rows = _sheet_rows(wb, 'PORTFOLIO_KPI')
    row = _find(rows, 'Alpha Co')
    assert row is not None
    rev, eb = row[6], row[7]                               # Revenue, EBITDA columns
    # REDDENING: held revenue must render ⚠HELD and MUST NOT leak its untrusted 99
    assert isinstance(rev, str) and 'HELD' in rev and '99' not in rev
    # a confirmed sibling still shows its real number (proves it's not blanket-hiding)
    assert eb == 5.0


# 3 ── HONEST NAV HOLD + reddening: NAV held ⇒ TVPI identity HELD, never PASS ──
def test_nav_held_makes_tvpi_identity_hold_REDDENING():
    wb = mw.build_master(_base_cir(nav_held=True), files=['x'])
    rows = _sheet_rows(wb, 'MOIC_TVPI_DPI')
    tvpi_metric = _find(rows, 'TVPI ((Dist + NAV) / Called)')
    identity = _find(rows, 'TVPI = DPI + RVPI')
    assert tvpi_metric is not None and 'HELD' in str(tvpi_metric[1])   # metric held
    assert identity is not None and 'HELD' in identity and 'PASS' not in identity
    # DPI still computes (called+distributed present) — the hold is scoped to NAV, not blanket
    dpi = _find(rows, 'DPI (Distributed / Called)')
    assert dpi is not None and isinstance(dpi[1], float)


def test_nav_present_makes_tvpi_identity_pass_POSITIVE():
    # the same fixture with NAV present must go PASS — proves #3 isn't green-by-accident
    wb = mw.build_master(_base_cir(nav_held=False), files=['x'])
    identity = _find(_sheet_rows(wb, 'MOIC_TVPI_DPI'), 'TVPI = DPI + RVPI')
    assert identity is not None and 'PASS' in identity and 'HELD' not in identity


# 4 ── DETERMINISM (correctly scoped) ────────────────────────────────────────
# Determinism = same CIR → same DATA. Raw xlsx bytes embed WALL-CLOCK in two container
# layers (docProps/core.xml + every ZIP member's date_time) that openpyxl stamps at WRITE
# time — that is when the file was written, not what it contains. So we assert (a) cell
# values are identical (the guarantee), and (b) the byte-REPRODUCIBLE saver (which strips
# both wall-clock layers) yields identical files — the artifact-level guarantee, robust by
# construction (no time gap can perturb it), NOT the flaky raw-bytes compare that was the F.
def _all_cells(wb):
    return {name: _sheet_rows(wb, name) for name in wb.sheetnames}


def test_build_master_cell_values_are_deterministic():
    a = mw.build_master(_base_cir(), files=['x'])
    b = mw.build_master(_base_cir(), files=['x'])
    assert _all_cells(a) == _all_cells(b), 'build_master cell values are not a pure function of the CIR'


def test_save_reproducible_is_byte_identical(tmp_path):
    pa, pb = tmp_path / 'a.xlsx', tmp_path / 'b.xlsx'
    mw.save_reproducible(mw.build_master(_base_cir(), files=['x']), str(pa))
    mw.save_reproducible(mw.build_master(_base_cir(), files=['x']), str(pb))
    assert pa.read_bytes() == pb.read_bytes(), 'reproducible save is not byte-identical'
    # REDDENING: a raw openpyxl save is NOT byte-stable (proves the normaliser is load-bearing,
    # not that xlsx happens to be stable) — the two raw saves differ once wall-clock advances.
    import time as _t
    raw1 = io.BytesIO(); mw.build_master(_base_cir(), files=['x']).save(raw1)
    _t.sleep(1.1)
    raw2 = io.BytesIO(); mw.build_master(_base_cir(), files=['x']).save(raw2)
    assert raw1.getvalue() != raw2.getvalue(), 'raw save unexpectedly stable — reddening control void'


# 5 ── STALENESS: tunable threshold + fail-closed on an unparseable date ──────
def test_staleness_flag_tunable_and_failclosed():
    fund_ym = mw._parse_ym('2026-06-30')
    assert mw._months_old('Dec-23 Actual', fund_ym) == 30      # old → will flag
    assert mw._months_old("Feb'26", fund_ym) == 4              # recent → won't flag
    assert mw._months_old('Q4', fund_ym) is None               # unparseable → NEVER guessed stale
    assert mw._months_old('', fund_ym) is None
    # the KPI grid flags a stale company and not a fresh one, at the default threshold
    cir = _base_cir()
    cir.records.append(Record('mis', entity_id='Alpha Co', fields={
        'company': 'Alpha Co', 'revenue': _fig('revenue', 10, col='Dec-23 Actual'),
        'ebitda': _fig('ebitda', 2, col='Dec-23 Actual'),
        'cash': _fig('cash', 3, col='Dec-23 Actual'), 'headcount': _fig('headcount', 5, col='Dec-23 Actual')}))
    row = _find(_sheet_rows(mw.build_master(cir, files=['x']), 'PORTFOLIO_KPI'), 'Alpha Co')
    assert row is not None and mw.STALE in str(row[5])         # Stale? column flagged


# 6 ── SLOW real-file end-to-end (June as-of): 13 sheets, checks PASS, stale ──
IN = 'backend/media/preingest/trivesta/100e86d5/in'


@pytest.mark.slow
@pytest.mark.skipif(not os.path.isdir(IN), reason='real fixture files not present')
def test_faithful_workbook_on_real_files():
    import glob
    from backend.dataimport.preingest3 import pipeline
    from backend.dataimport.preingest3.ratecard import default_inr_card
    files = [(os.path.basename(p), p) for p in sorted(glob.glob(os.path.join(IN, '*.xlsx')))]
    res = pipeline.run(files, as_of='2026-06-30', org='default', rate_card=default_inr_card('2026-06-30'))
    cir = res.cir if hasattr(res, 'cir') else res
    wb = mw.build_master(cir, rate_card=default_inr_card('2026-06-30'), files=[f for f, _ in files])
    assert wb.sheetnames == mw._SHEETS                          # exactly the 13-sheet shape, in order
    m = _sheet_rows(wb, 'MOIC_TVPI_DPI')
    assert 'PASS' in _find(m, 'Total FV = Σ company FV')        # 827 = 827
    assert 'PASS' in _find(m, 'TVPI = DPI + RVPI')              # coherent June as-of
    # Analisa (Dec-2023 data) is flagged stale in the KPI grid
    kpi = _sheet_rows(wb, 'PORTFOLIO_KPI')
    analisa = next((r for r in kpi if r and isinstance(r[0], str) and r[0].startswith('Analisa')), None)
    assert analisa is not None and mw.STALE in str(analisa[5])
    # SECTOR_ALLOCATION is populated (not the empty v3 state)
    assert len(_sheet_rows(wb, 'SECTOR_ALLOCATION')) > 4


# 7 ── NAV §5.1 two-path build + reddening: both paths shown, not one echoed ──
def _cir_with_bs():
    """base CIR + the §5.1 balance-sheet fields on the NAV record (ΣFV 300 of the 2 invs)."""
    cir = _base_cir()
    nav = next(r for r in cir.records if r.domain == 'nav')
    nav.fields['bs_components'] = [('Cash & cash equivalents', '22', 'fund ac and PnL!F6'),
                                   ('Receivables (interest/div)', '6', 'fund ac and PnL!F7'),
                                   ('Mgmt fee payable', '4', 'fund ac and PnL!F8'),
                                   ('Other liabilities (borrowings)', '8', 'fund ac and PnL!F9')]
    nav.fields['bs_portfolio_fv'] = '300'          # = Σ FV (100 + 200)
    nav.fields['bs_gross_nav'] = '316'             # 300 + (22+6) − (4+8)
    nav.fields['bs_lp_nav'] = '254.4'              # 316 − carry 61.6
    cir.checks.append({'id': 'nav_two_path_reconciliation', 'class': 'soft', 'status': 'pass',
                       'lhs': '867.85', 'rhs': '316.00',
                       'detail': 'gross NAV: roll-forward 867.85 vs §5.1 316.00; Δ=551.85'})
    return cir


def test_nav_two_path_51_build_renders_both_paths_REDDENING():
    rows = _sheet_rows(mw.build_master(_cir_with_bs(), files=['x']), 'NAV_CALC')
    flat = [str(c) for r in rows for c in r]
    # Path A (§5.1) rendered: the FV line, cash as (+), fee-payable as (−)
    assert any(c.startswith('(+) Investments at fair value') for c in flat)
    assert any(c.startswith('(+) Cash') for c in flat)
    assert any(c.startswith('(−) Mgmt fee payable') for c in flat)
    # BOTH gross values present and DIFFERENT — proves two independent paths, not one echoed
    gross_vals = [r for r in rows if r and str(r[0]).startswith('(=) Gross NAV')]
    shown = {r[1] for r in gross_vals}
    assert 316.0 in shown and 867.85 in shown, f'expected both §5.1 316 and roll-forward 867.85, got {shown}'
    # the two-path reconciliation row is materialised
    assert _find(rows, '§5.1 vs roll-forward') is not None


def test_nav_no_bs_components_falls_back_to_rollforward_only_POSITIVE():
    # no §5.1 fields → Path A must NOT be drawn (proves Path A is data-gated, not always-on)
    rows = _sheet_rows(mw.build_master(_base_cir(), files=['x']), 'NAV_CALC')
    flat = [str(c) for r in rows for c in r]
    assert not any(c.startswith('(+) Investments at fair value') for c in flat)
    assert any('PATH B' in c for c in flat)        # roll-forward still there


# 8 ── WATERFALL: crystallised 0 carry (pre-capital-back) + illustrative band ──
def _cir_with_terms():
    from backend.dataimport.preingest3.fund_terms import FundTerm
    cir = _base_cir()          # nav gross 867.85 · carry 61.6 · ceiling 84.8 · called 600 · dist 70

    def _t(concept, val):
        return FundTerm(concept=concept, value=Decimal(str(val)), unit='percent', base='n/a',
                        phase='whole_life', text_value=None, provenance=_pv())
    cir.records.append(Record('fund_terms', entity_id='fund', fields={
        'carried_interest': _t('carried_interest', '0.20'), 'hurdle_rate': _t('hurdle_rate', '0.08'),
        'catch_up': _t('catch_up', '1.0')}))
    return cir


def test_waterfall_crystallised_zero_carry_and_illustrative_band_REDDENING():
    rows = _sheet_rows(mw.build_master(_cir_with_terms(), files=['x']), 'WATERFALL_EUR')
    # crystallised: 70 all return-of-capital; GP carry crystallised EXACTLY 0 below capital-back
    roc = _find(rows, '→ Return of capital paid')
    assert roc is not None and roc[1] == 70.0
    carry0 = _find(rows, '→ GP carry crystallised')
    # REDDENING: if the waterfall ever credited carry before LPs are capital-back this is != 0
    assert carry0 is not None and carry0[1] == 0.0
    # illustrative implied GP carry = 20% × profit 337.85 (full catch-up) = 67.57
    impl = _find(rows, '(=) Implied GP carry')
    assert impl is not None and abs(impl[1] - 67.57) < 0.05
    # cross-check band present and PASS (61.6 ≤ 67.57 ≤ 84.8)
    band = _find(rows, 'cross-check: accrued')
    assert band is not None and 'PASS' in band


# 9 ── canonical percent formatter (ONE impl for nav/_fmt_pct/_pctlabel) + reddening ──
def test_format_pct_canonical_with_negative_controls():
    # positive — the shapes every call site needs, all through the single primitive
    assert format_pct(Decimal('1')) == '100%'                     # integer-valued: zeros kept
    assert format_pct(Decimal('0.02')) == '2%'
    assert format_pct(Decimal('0.2')) == '20%'
    assert format_pct(Decimal('0.025')) == '2.5%'
    assert format_pct(Decimal(70) / Decimal(600)) == '11.67%'     # computed fraction: rounded
    assert format_pct(None) == 'n/r'
    # REDDENING 1 (the 100→1 bug) — bare rstrip('0') with NO decimal guard eats real zeros;
    # prove the naive transform corrupts, and the canonical does NOT reproduce it
    naive_100 = format(Decimal('1') * 100, 'f').rstrip('0').rstrip('.')
    assert naive_100 == '1', 'negative control void: naive strip should corrupt 100→1'
    assert format_pct(Decimal('1')) != naive_100 + '%'
    # REDDENING 2 (the 11.666… bug) — without quantize a computed fraction keeps the full tail
    unrounded = format(Decimal(70) / Decimal(600) * 100, 'f')
    assert unrounded.startswith('11.6666'), 'negative control void: unrounded should be a long tail'
    assert format_pct(Decimal(70) / Decimal(600)) != unrounded + '%'


# 10 ── RECONCILIATION: a PASSING soft check is published (U7), not dropped ───
def test_reconciliation_publishes_passing_soft_check_REDDENING():
    cir = _base_cir()
    cir.checks.append({'id': 'portfolio_revenue_vs_fund', 'class': 'soft', 'status': 'pass',
                       'lhs': '100', 'rhs': '113', 'detail': '13% variance within tolerance — still published'})
    cir.checks.append({'id': 'tranches_sum_to_cost', 'class': 'hard', 'status': 'pass',
                       'lhs': '448', 'rhs': '448', 'detail': 'ties'})
    rows = _sheet_rows(mw.build_master(cir, files=['x']), 'RECONCILIATION')
    # REDDENING: a PASSING SOFT check must appear — if the sheet filtered to failures it vanishes
    soft = _find(rows, 'portfolio_revenue_vs_fund')
    assert soft is not None and soft[1] == 'soft' and 'PASS' in soft
    assert _find(rows, 'tranches_sum_to_cost') is not None      # hard pass also published
    assert any('TOTALS' in str(r[0]) for r in rows)


# 11 ── DASHBOARD_BRIDGE is a widget→source map (not the flat dump) ───────────
def test_dashboard_bridge_is_widget_map_with_granular_drilldown():
    rows = _sheet_rows(mw.build_master(_base_cir(), files=['x']), 'DASHBOARD_BRIDGE')
    flat = [str(c) for r in rows for c in r]
    assert any('Widget' in c for c in flat) and any('Calc Logic' in c for c in flat)   # widget-map header
    assert any(r and r[0] == 'Fund Overview' for r in rows)                            # a curated widget row
    assert any('GRANULAR PROVENANCE FEED' in c for c in flat)                          # drill-down retained


# 12 ── CURRENCY UNIT on every money header (attestation) + reddening ─────────
def test_hdr_appends_cr_only_on_exact_money_names_REDDENING():
    # every ₹Cr money value must carry its unit — but EXACT-match only. The reddening core: if _hdr ever
    # switched to substring matching, look-alikes ('EBITDA Margin %', 'EBITDA basis/adj') would be
    # mislabelled as ₹Cr — a wrong unit on a %/text column. This test goes RED against exactly that.
    out = mw._hdr(['Revenue', 'EBITDA', 'EBITDA (Standard)', 'Cash', 'Cost', 'Fair value',
                   'EBITDA Margin %', 'EBITDA basis/adj', 'Head-count', 'MOIC', 'IRR', 'Gain %',
                   'Company', 'Basis', 'Val Method'])
    assert out[:6] == ['Revenue (₹Cr)', 'EBITDA (₹Cr)', 'EBITDA (Standard) (₹Cr)',
                       'Cash (₹Cr)', 'Cost (₹Cr)', 'Fair value (₹Cr)']            # must-handle
    for bare in ['EBITDA Margin %', 'EBITDA basis/adj', 'Head-count', 'MOIC', 'IRR', 'Gain %',
                 'Company', 'Basis', 'Val Method']:                               # must-not-misfire
        assert bare in out and f'{bare} ({mw._CCY})' not in out


def test_money_headers_carry_unit_in_built_workbook():
    cir = _base_cir()
    cir.records.append(Record('mis', entity_id='Alpha Co', fields={
        'company': 'Alpha Co', 'revenue': _fig('revenue', 12, col='2026-02-28'),
        'ebitda': _fig('ebitda', 3, col='2026-02-28'), 'cash': _fig('cash', 8, col='2026-02-28'),
        'headcount': _fig('headcount', 40, col='2026-02-28')}))
    wb = mw.build_master(cir, files=['x'])
    def hdrs(name):
        return [c for row in _sheet_rows(wb, name) for c in row if isinstance(c, str)]
    kpi = hdrs('PORTFOLIO_KPI')
    for h in ('Revenue', 'EBITDA', 'Cash'):
        assert f'{h} ({mw._CCY})' in kpi, f'PORTFOLIO_KPI must label {h} with the unit'
    assert 'EBITDA Margin %' in kpi and f'EBITDA Margin % ({mw._CCY})' not in kpi   # % stays bare
    assert 'Head-count' in kpi and f'Head-count ({mw._CCY})' not in kpi            # count stays bare
    assert any(c.startswith('PORTFOLIO KPI TRACKER') and mw._CCY in c for c in kpi)  # banner states ₹Cr
    pm = hdrs('PORTFOLIO_MASTER')
    assert f'EBITDA (Standard) ({mw._CCY})' in pm and f'Monthly Burn ({mw._CCY})' in pm
    assert 'Equity %' in pm and f'Equity % ({mw._CCY})' not in pm
    val = hdrs('VALUATIONS')
    assert f'Fair Value ({mw._CCY})' in val and f'Unrealised Gain ({mw._CCY})' in val
    for r in ('Gain %', 'IRR', 'Multiple'):                                         # ratios stay bare
        assert r in val and f'{r} ({mw._CCY})' not in val
    moic = hdrs('MOIC_TVPI_DPI')
    assert f'Cost ({mw._CCY})' in moic and f'Fair value ({mw._CCY})' in moic
    assert 'Value' in moic and f'Value ({mw._CCY})' not in moic                     # MOIC 'Value'=ratios
