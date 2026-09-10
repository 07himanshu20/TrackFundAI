"""TIER-S load generator (test-only) — provenance-tracked SCRAMBLED variants for the coverage/robustness
track (#5) of the 100-file readiness decomposition. Unlike Tier-V (volume replicas for rate/resource),
Tier-S changes BOTH the fingerprint AND the structural layout, and — because WE generate the
perturbation — it hands back a GROUND-TRUTH MAP (a known sentinel value planted next to a known concept
label, at a known post-transform cell). That is what makes a real locator-robustness test possible with
NO real diverse files and NO tokens for the instrument: a later locating run asserts located-value ==
planted-sentinel across every layout perturbation.

CRITICAL (item 4, direct consequence of: dedupe keys on content_fp, not raw bytes): every variant MUST
have a DISTINCT content_fp or the reuse cache/golden store collapses it and the robustness run is
vacuous. The unique sentinels guarantee it; `assert_distinct_content_fp` enforces it by construction.

Universality guard: the transforms are layout-AGNOSTIC (generic: reorder/rename sheets, inject noise
rows, relabel via a small generic synonym map, plant sentinels in any text->numeric row) — NOT tuned to
the 15 files. A variant the system fails on is a real universality gap to fix at the mechanism, never to
special-case. HONEST BOUNDARY: Tier-S tests robustness to transformations WE applied, not to novel
real-world messiness we did not imagine — a strong proxy, not a substitute for the real corpus.

The locating RUN (assert located == sentinel) is the only token spend and is DEFERRED to a greenlight;
this module + its self-test are token-free (fingerprints + re-read only)."""
from __future__ import annotations

import random
from dataclasses import dataclass
from typing import List

import openpyxl

from backend.dataimport.preingest3.identity import compute_identity

# A small GENERIC relabel map — a header-synonym transform, illustrative and not file-specific; it only
# fires where such a header happens to exist. The load-bearing transforms below are label-independent.
_SYNONYMS = {
    'Particulars': 'Line Item', 'Revenue': 'Total Income', 'EBITDA': 'Operating Profit',
    'Company': 'Entity', 'Cost': 'Book Value', 'Total': 'Aggregate',
}


@dataclass
class Planted:
    """Ground truth: a locator asked for `concept` on this variant must return `sentinel` (which lives
    at `sheet`!`cell` after the layout scramble)."""
    concept: str
    sentinel: float
    sheet: str
    cell: str


def scramble(src: str, dst: str, *, seed: int) -> List[Planted]:
    """Produce ONE scrambled variant of `src` at `dst` and return its ground-truth map. Transforms
    (all generic): header relabel, noise-row injection, sheet rename, sheet reorder, then sentinel
    planting (recorded AFTER the layout shifts so cell addresses are accurate)."""
    rng = random.Random(seed)
    wb = openpyxl.load_workbook(src)

    for ws in wb.worksheets:                                   # 1. generic header/label relabel
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.strip() in _SYNONYMS:
                    c.value = _SYNONYMS[c.value.strip()]

    for ws in wb.worksheets:                                   # 2. inject noise rows (shifts layout down)
        ws.insert_rows(1, amount=1 + rng.randint(1, 3))
        ws['A1'] = f'-- tier_s noise seed={seed} --'

    for ws in wb.worksheets:                                   # 3. rename sheets (bounded to 31 chars)
        ws.title = (ws.title[:24] + f'_s{seed}')[:31]

    order = list(wb._sheets)                                   # 4. reorder sheets
    rng.shuffle(order)
    wb._sheets = order

    planted: List[Planted] = []                               # 5. plant sentinels in text->numeric rows
    uid = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            cells = list(row)
            for i in range(len(cells) - 1):
                lbl, val = cells[i].value, cells[i + 1].value
                if isinstance(lbl, str) and lbl.strip() and isinstance(val, (int, float)) and not isinstance(val, bool):
                    uid += 1
                    sentinel = float(900_000_000 + seed * 100_000 + uid)   # unique + recognizable
                    cells[i + 1].value = sentinel
                    planted.append(Planted(lbl.strip(), sentinel, ws.title, cells[i + 1].coordinate))
                    break                                     # one sentinel per row
    wb.save(dst)
    return planted


def assert_distinct_content_fp(paths: List[str]) -> None:
    """Item-4 discipline, enforced by construction: every variant must have a DISTINCT content_fp
    (else dedupe collapses it). Reddens if two variants would be treated as the same file."""
    fps = {}
    for p in paths:
        fp = compute_identity(p, p).content_fp
        assert fp not in fps, f'DUPLICATE content_fp: {p} collides with {fps[fp]} — dedupe would collapse it'
        fps[fp] = p


def verify_ground_truth(dst: str, planted: List[Planted]) -> None:
    """Map integrity (token-free): re-read the saved variant and confirm each planted sentinel is
    actually present at its recorded sheet!cell. Proves the ground-truth map a locating run will assert
    against is internally correct BEFORE any tokens are spent."""
    wb = openpyxl.load_workbook(dst, data_only=True)
    for p in planted:
        got = wb[p.sheet][p.cell].value
        assert got == p.sentinel, f'ground-truth broken: {p.sheet}!{p.cell} holds {got!r}, expected {p.sentinel}'
