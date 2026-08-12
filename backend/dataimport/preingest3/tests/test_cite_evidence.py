"""CITE-EVIDENCE META-TEST (advisor Step 3) — the ruler every net test inherits.

The architecture's promise is that a figure carries the evidence that reconstructs it.
This test enforces that promise UNIVERSALLY: for every emitted figure on every MIS file
(production config), the provenance must let you rebuild the number from raw cells —
"cite the cells + operation that reconstruct the value."

Four assertions per emit:
  1. It cites evidence at all — derived_from is non-empty.
  2. The provenance model is structurally honest — cell=='' IFF the value is a multi-cell
     sum. This is the assertion that kills the misleading case the advisor named: a TTM
     figure (Σ of 12 monthly cells) that cites a single cell would look verifiable and
     reconstruct to the wrong number.
  3. Every cited cell is actually numeric in the raw sheet (no dangling citation).
  4. RECONSTRUCTION: applying the basis operation to the cited raw cells reproduces the
     value. A count must EQUAL its single cited cell exactly. A money figure's value_cr
     must be Σ(cited raw cells) × a single RECOGNISED unit scale (INR-only prod: 1 [₹Cr],
     0.1 [millions], 0.01 [lakhs], 1e-4 [thousands], 1e-7 [absolute ₹]). A wrong-cell or
     wrong-count citation makes the implied scale land between recognised factors and fails.

SCOPE — PROVENANCE-INTEGRITY, NOT VALUE-CORRECTNESS (advisor 2026-07-26, do not conflate):
this ruler proves the value is a scaled sum of the cited cells — i.e. the provenance is
HONEST and the cells are the right cells. It does NOT prove the SCALE is right: the money
check searches {1,.1,.01,1e-4,1e-7} for ANY factor that fits, so a WRONG scale that is
internally consistent with the cells passes green while the number is wrong. Value/scale
correctness comes ONLY from the value-audit against EXTERNAL ground truth (test_coverage_gate
GT + the slot-by-slot audits at the LLM-on and U6-wiring events). Cite-evidence is NECESSARY,
NOT SUFFICIENT: every NEW emit (net-sourced, LLM-filled, un-masked foreign-ccy, fund figures)
still needs an independent ground-truth value-audit — this test cannot catch a self-consistent
scale error. Two rulers, named distinctly; never let the cheap one stand in for the expensive one.

Slow: extracts all 10 MIS files. Skipped if fixtures absent. When the net + LLM emit new
figures, they inherit this ruler automatically — no per-file provenance test needed.
"""
import os
from decimal import Decimal

import openpyxl
import pytest

from backend.dataimport.preingest3 import fund_anchor
from backend.dataimport.preingest3.extract import extract_company, MIS_CONCEPTS
from backend.dataimport.preingest3.ratecard import default_inr_card
from backend.dataimport.preingest3.contract import concept_measure
from backend.dataimport.preingest3.cir import Figure

IN = 'backend/media/preingest/trivesta/100e86d5/in'
pytestmark = [
    pytest.mark.slow,
    pytest.mark.skipif(not os.path.isdir(IN), reason='real fixture files not present'),
]

FUND = ['TFAI_Investments_and_Deployment.xlsx', 'TFAI_Valuations_and_Exits_Q2FY26.xlsx']
MIS = [
    ('Hubler', 'AVF_2026_03_11_P_Hubler_MIS_Feb26.xlsx', 'hubbler'),
    ('Agnikul', 'AVF_2026_03_12_P_Agnikul_MIS_Feb26.xlsx', 'agnikul'),
    ('Clientell', 'AVF_2026_03_16_P_Clientell_MIS_Feb26.xlsx', 'clientell'),
    ('InstaAstro', 'AVF_2026_03_18_P_InstaAstro_MIS_Feb_2026.xlsx', 'instaastro'),
    ('LDC', 'AVF_2026_03_20_P_LDC_Detailed_MIS_Feb26.xlsx', 'ldc'),
    ('Aliste', 'AVF_2026_03_26_P_Aliste_MIS_Feb26.xlsx', 'aliste'),
    ('CPC', 'CPC_Monthly_MIS-_May_25_to_be_sent_to_EL.xlsx', 'cpc'),
    ('CPM', '0625_CPM_Monthly_Report-R1.xlsx', 'chemopharm'),
    ('CSS', '0525_CSS_Monthy_Report_-_Consol_Updated.xlsx', 'css'),
    ('Analisa', '01_Monthly_Financial_Presentation_2025_May_Analisa.xlsx', 'analisa'),
]
# native→₹Cr scales possible under the INR-only shipping card (FX-scaled emits HOLD).
_SCALES = [Decimal('1'), Decimal('0.1'), Decimal('0.01'), Decimal('0.0001'), Decimal('0.0000001')]


def _raw_sum(wb_cache, path, sheet, cells):
    ws = wb_cache.get((path, sheet))
    if ws is None:
        ws = openpyxl.load_workbook(path, data_only=True, read_only=True)[sheet]
        wb_cache[(path, sheet)] = list(ws.iter_rows(values_only=True))
        ws = wb_cache[(path, sheet)]
    total, raws = Decimal('0'), []
    for ref in cells:
        r, c = openpyxl.utils.coordinate_to_tuple(ref)   # 1-based (row, col)
        v = ws[r - 1][c - 1] if r - 1 < len(ws) and c - 1 < len(ws[r - 1]) else None
        assert isinstance(v, (int, float)), f'cited cell {sheet}!{ref} is not numeric: {v!r}'
        total += Decimal(str(v))
        raws.append(Decimal(str(v)))
    return total, raws


def _recognised_scale(implied: Decimal) -> bool:
    return any(abs(implied - s) <= s * Decimal('0.02') for s in _SCALES)


def test_every_emit_reconstructs_from_its_cited_cells():
    anchors = fund_anchor.build_fund_anchors([os.path.join(IN, f) for f in FUND])
    wb_cache: dict = {}
    checked, failures = 0, []
    for name, fn, token in MIS:
        ca = next((c for k, c in anchors.items() if token in k.lower()), None)
        path = os.path.join(IN, fn)
        rec = extract_company(name, path, rate_card=default_inr_card('2026-02-28'), entity=name,
                              domicile=(ca.domicile if ca else None),
                              anchor_cr=(ca.anchor_cr if ca else None), use_model=False)
        figs = {v.concept: v for v in rec.fields.values() if isinstance(v, Figure)}
        for c in MIS_CONCEPTS:
            f = figs.get(c)
            if f is None or f.value_cr is None:
                continue                                  # held/gap → nothing to reconstruct
            p = f.provenance
            tag = f'{name}/{c}'
            # (1) cites evidence
            if not p.derived_from:
                failures.append(f'{tag}: emits {f.value_cr} but cites NO cells'); continue
            # (2) structurally honest: single value cell ⟺ point figure; '' ⟺ multi-cell sum
            is_sum = len(p.derived_from) > 1
            if (p.cell == '') != is_sum:
                failures.append(f'{tag}: cell={p.cell!r} but {len(p.derived_from)} cells (sum must cite "")')
                continue
            # (3)+(4) cells are numeric AND reconstruct the value
            try:
                total, raws = _raw_sum(wb_cache, path, p.sheet, p.derived_from)
            except AssertionError as e:
                failures.append(f'{tag}: {e}'); continue
            if concept_measure(c) == 'count':
                if not (len(raws) == 1 and Decimal(f.value_cr) == raws[0]):
                    failures.append(f'{tag}: count {f.value_cr} != single cited cell {raws}')
            else:                                          # money
                if total == 0 or (total > 0) != (f.value_cr > 0):
                    failures.append(f'{tag}: Σraw {total} sign-mismatches value {f.value_cr}')
                elif not _recognised_scale(abs(f.value_cr) / abs(total)):
                    failures.append(f'{tag}: implied scale {abs(f.value_cr)/abs(total)} '
                                    f'(value {f.value_cr} / Σraw {total}) is not a recognised unit — '
                                    f'wrong cells or wrong count cited')
            checked += 1
    assert checked >= 10, f'expected ≥10 emits to reconstruct, only saw {checked}'
    assert not failures, 'cite-evidence reconstruction failures:\n  ' + '\n  '.join(failures)
