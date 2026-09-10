"""
U3 — Three-level file identity.

Closes D3. A raw-byte hash answers only "is this the identical file?" — too
narrow, because opening a workbook in Excel and pressing save rewrites bytes
while changing nothing that matters, which under a byte cache re-invoked the
model and broke determinism. Three questions, three fingerprints (Section 07):

  L1 BYTE    — hash of the file as uploaded. Fast path only, never the sole key.
  L2 CONTENT — hash of sheet names + cell addresses + cell VALUES, canonically
               ordered. Survives re-saves / metadata edits. The golden-record
               key (delivers the determinism guarantee G1).
  L3 LAYOUT  — hash of structure with all VALUES stripped: sheet names, header
               text, label-column text. Identical for the same template in a
               different month, or a different company on the same template.
               The Layout Template Registry key (delivers the cost guarantee G4).

Canonical ordering: files are sorted by content fingerprint before processing —
never by upload order or filename — so two users submitting the same set in a
different sequence obtain the same workbook.
"""
from __future__ import annotations

import datetime as _dt
import hashlib
import os
import xml.etree.ElementTree as _ET
import zipfile
from collections import OrderedDict
from dataclasses import dataclass
from typing import List, Optional

import openpyxl
from openpyxl.utils import get_column_letter as _col_letter

# The cache logic-version is a hash of the library's source, so ANY change to value logic
# auto-invalidates cached results (Inc-5). Preferred to a manually-bumped version string
# (which rests on remembering to bump it — forget once and a cache hit serves a pre-fix
# result, the same silent-staleness class as a content_fp-only key). It hashes the WHOLE
# package, NOT a hand-list of "value modules": tracing extract_company's real call graph
# proved a hand-list (extract/family/tiers/reconcile/periods/contract/units/quantity) SILENTLY
# MISSED 6 value-affecting modules — gate (the binding decision), lexicon (concept matching),
# profiler (grid/axis), ratecard (FX conversion code), cir, identity. Enumeration-from-memory
# is exactly the completeness hole this hash exists to close, so we don't enumerate: hashing
# every module means over-inclusion (a harmless false cache-miss on a non-value edit) but NEVER
# under-inclusion (a stale-logic hole). Also covers the use_model=True model path (locator/llm/
# namematch/templates) that a deterministic trace can't see but the cache serves under one key.
_LOGIC_VERSION: Optional[str] = None

# byte_fp → (content_fp, layout_fp). compute_identity is a PURE function of file bytes but was
# called 5-7× per file per run (canonical-order sort + per-file loops + per-extract), each time
# re-parsing the whole workbook with openpyxl (~6s/file) — the dominant cost of a model-OFF run.
# Memoising by byte_fp (the exact-bytes hash) skips the re-parse on every repeat while staying
# provably safe: any edit changes the bytes → new byte_fp → miss → recompute (the edit-retest
# invariant). Content-derived, so label/path (call-specific) are NOT part of the key.
# BOUNDED (LRU): a plain dict would grow without limit in a long-lived client process (a slow
# memory leak — invisible in the 15-file dev suite, real in production). The cap turns it into an
# LRU: past _IDENTITY_CACHE_MAX distinct files the least-recently-used entry is evicted. Eviction
# is VALUE-NEUTRAL — an evicted file simply misses and recomputes the identical (pure) fingerprints,
# never a stale/wrong number. Entries are tiny (three hex strings), so the cap is generous.
_IDENTITY_CACHE_MAX = 1024
_IDENTITY_CACHE: "OrderedDict[str, tuple]" = OrderedDict()


def clear_identity_cache() -> None:
    """Drop the in-process identity memo (for tests/long-lived processes; not needed for
    correctness — the cache is byte-keyed and pure)."""
    _IDENTITY_CACHE.clear()


def _cache_put(byte_fp: str, content_fp: str, layout_fp: str) -> None:
    """Insert/refresh an identity entry and evict the least-recently-used beyond the cap, so the
    memo cannot grow without bound in a long-lived process. Value-neutral: an evicted file just
    misses next time and recomputes the identical (pure) fingerprints — never a stale/wrong one."""
    _IDENTITY_CACHE[byte_fp] = (content_fp, layout_fp)
    _IDENTITY_CACHE.move_to_end(byte_fp)
    while len(_IDENTITY_CACHE) > _IDENTITY_CACHE_MAX:
        _IDENTITY_CACHE.popitem(last=False)


# PARSE-ONCE: the per-run cache of the single unified workbook parse (fingerprints + grid), keyed by
# (path, mtime, size) like the profiler cache. compute_identity AND profiler.profile_file both read it,
# so a workbook is opened ONCE per run instead of twice (the identity fingerprint pass and the profiler
# grid pass used to be separate openpyxl loads — the dominant sequential cost at scale). It holds grids,
# so it is CLEARED at each pipeline.run start (clear_parse_cache); the tiny fingerprint cache above
# persists across runs (its entries are immutable strings).
_PARSE_CACHE: dict = {}


def clear_parse_cache() -> None:
    """Drop the per-run unified-parse cache (grids). Called at each pipeline.run start so a run reuses
    each file's single parse but no run holds another run's grids. Memory-bounding + cross-run
    freshness; correctness-neutral (a re-parse yields an identical result). Also resets the per-run
    calamine→openpyxl fallback log so each run reports its own fallback rate (G-FALLBACK)."""
    _PARSE_CACHE.clear()
    _FALLBACK_LOG.clear()


@dataclass(frozen=True)
class FileIdentity:
    label: str
    path: str
    byte_fp: str
    content_fp: str
    layout_fp: str
    currencies: tuple = ()      # distinct currency hints seen (for Rate Card coverage)


def _sha(parts) -> str:
    h = hashlib.sha256()
    for p in parts:
        h.update(p.encode('utf-8', 'replace'))
        h.update(b'\x1f')
    return h.hexdigest()


def _is_volatile_value(v) -> bool:
    """A VALUE that changes reporting-period to reporting-period — numbers AND
    dates/times. These are stripped from the LAYOUT fingerprint (else the same
    template next month, with new numbers and a new period-end date, would hash
    differently and be wrongly treated as a new layout — defeating U3/G4). They
    are kept in the CONTENT fingerprint, which is meant to differ per period."""
    if isinstance(v, bool):
        return False
    return isinstance(v, (int, float, _dt.datetime, _dt.date, _dt.time))


def _parse_workbook_openpyxl(path: str):
    """The openpyxl read-only, data_only pass — the BYTE-EXACT REFERENCE reader and the whole-file
    fallback for calamine (see _parse_workbook). In one iteration it accumulates the content/layout
    fingerprint parts (identical, line-for-line, to the historical compute_identity loop) AND the
    per-sheet value grid (identical to the historical profiler._read_grid: [[cell.value ...] ...]) —
    so unifying the two parses cannot change either output. Returns (byte_fp, content_fp, layout_fp,
    grid, error)."""
    with open(path, 'rb') as fh:
        byte_fp = hashlib.sha256(fh.read()).hexdigest()
    content_parts: List[str] = []
    layout_parts: List[str] = []
    grid: dict = {}
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        # Non-workbook or unreadable → identity still defined by bytes; the profiler re-raises
        # this to reproduce its exact error-profile. content/layout fall back to the byte hash.
        return (byte_fp, byte_fp, byte_fp, {}, e)

    for sn in wb.sheetnames:                       # sheetnames order is canonical
        ws = wb[sn]
        content_parts.append(f'#SHEET#{sn}')
        layout_parts.append(f'#SHEET#{sn}')
        rows_out: list = []
        for row in ws.iter_rows():
            vals = []
            for cell in row:
                v = cell.value
                vals.append(v)                      # grid keeps EVERY cell, positionally (profiler parity)
                if v is None or v == '':
                    continue
                coord = cell.coordinate
                content_parts.append(f'{coord}={v!r}')
                if not _is_volatile_value(v):       # text = structure; values stripped
                    layout_parts.append(f'{coord}={str(v).strip().lower()!r}')
            rows_out.append(vals)
        grid[sn] = rows_out
    try:
        wb.close()
    except Exception:
        pass

    content_fp = _sha(content_parts) if content_parts else byte_fp
    layout_fp = _sha(layout_parts) if layout_parts else byte_fp
    return (byte_fp, content_fp, layout_fp, grid, None)


# ── Calamine fast reader (~18× openpyxl) + format-based OOR-date safety guard ────────────────────
# The single workbook parse dominates a model-off run (~91% of wall time). python-calamine (Rust) reads
# the same values ~18× faster, so it is the PRIMARY reader; openpyxl above stays the byte-exact reference
# and whole-file fallback. calamine is value-faithful after four normalizations proven byte-identical on
# the corpus (see _norm_calamine_value) EXCEPT one class it cannot represent: a date-FORMATTED cell whose
# serial is out of Excel's date range. openpyxl range-checks these and returns an error; calamine returns
# the raw serial (a ~2e9 number) which, extracted, would be a WRONG NUMBER (the cardinal sin). calamine
# cannot expose the number format, and the value alone is indistinguishable from a legitimate large figure
# (tens of thousands of legit cells exceed the max serial on the corpus). So the guard reads the number
# formats directly (styles.xml) and, if ANY date-formatted cell holds an out-of-range serial, the WHOLE
# file falls back to openpyxl — never a wrong number, never two readers mixed in one file. Fail-closed:
# any calamine/guard error → openpyxl. Every fallback is logged (G-FALLBACK: a high fallback rate would
# make the speedup illusory and must surface, not degrade silently).
try:
    from python_calamine import CalamineWorkbook as _CalamineWorkbook
    _HAS_CALAMINE = True
except Exception:                               # dependency missing → openpyxl-only, behaviour unchanged
    _HAS_CALAMINE = False

_MAX_XL_DATE_SERIAL = 2958465                    # 9999-12-31; openpyxl treats larger serials as errors
_SML_NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
_FALLBACK_LOG: list = []                         # [(path, reason)] this run — G-FALLBACK instrumentation


def clear_fallback_log() -> None:
    _FALLBACK_LOG.clear()


def fallback_log() -> list:
    """Files this run that fell back from calamine to openpyxl, each (path, reason) (G-FALLBACK)."""
    return list(_FALLBACK_LOG)


def _norm_calamine_value(v):
    """Map a calamine cell value to openpyxl's EXACT representation (each rule proven byte-identical on
    the corpus): ''→None (empty cell); integral float→int (openpyxl casts int iff the stored number has
    no decimal point — the corpus has ZERO float-but-integral counterexamples); date→datetime at midnight
    (openpyxl always returns datetime, never a bare date); and calamine's decoded '\\r' back to openpyxl's
    un-decoded '_x000D_' literal so multi-line strings hash identically. bool is preserved (never an int).
    NaN→None. OOR-date serials never reach here — such files fall back whole to openpyxl."""
    if v == '':
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, float):
        if v != v:                               # NaN → empty (never a value)
            return None
        return int(v) if v.is_integer() else v
    if isinstance(v, _dt.datetime):
        return v
    if isinstance(v, _dt.date):
        return _dt.datetime(v.year, v.month, v.day)
    if isinstance(v, str):
        return v.replace('\r', '_x000D_') if '\r' in v else v
    return v


def _date_format_style_ids(zf: zipfile.ZipFile) -> set:
    """Style indices (cellXfs order) whose number format is a DATE format — via openpyxl's own
    is_date_format so custom date codes are covered too. Bias to inclusion: a false positive only costs
    an extra openpyxl fallback (safe/slower); a false negative would let a wrong number through."""
    from openpyxl.styles.numbers import is_date_format, BUILTIN_FORMATS
    root = _ET.fromstring(zf.read('xl/styles.xml'))
    codes = dict(BUILTIN_FORMATS)
    for nf in root.iter(f'{_SML_NS}numFmt'):
        codes[int(nf.get('numFmtId'))] = nf.get('formatCode')
    date_ids = {i for i, c in codes.items() if c and is_date_format(c)}
    styles = set()
    cellXfs = root.find(f'{_SML_NS}cellXfs')
    if cellXfs is not None:
        for idx, xf in enumerate(cellXfs.findall(f'{_SML_NS}xf')):
            if int(xf.get('numFmtId', 0)) in date_ids:
                styles.add(idx)
    return styles


def _coord_to_ri_ci(ref: str) -> tuple:
    """'E264' → (263, 4): zero-based (row, col) from an A1 coordinate. Dependency-free so it can't drift
    with openpyxl's util surface."""
    i = 0
    while i < len(ref) and ref[i].isalpha():
        i += 1
    col = 0
    for ch in ref[:i]:
        col = col * 26 + (ord(ch) - 64)
    return int(ref[i:]) - 1, col - 1


def _sheet_order(zf: zipfile.ZipFile):
    """[(sheet_name, 'xl/worksheets/sheetN.xml'), …] in the workbook-declared order — the SAME order
    calamine's sheet_names and openpyxl's sheetnames use — so error cells can be overlaid onto calamine's
    grid by name. Maps each <sheet r:id> through workbook.xml.rels to its worksheet part."""
    wb = _ET.fromstring(zf.read('xl/workbook.xml'))
    rels = _ET.fromstring(zf.read('xl/_rels/workbook.xml.rels'))
    _RNS = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    rid_target = {r.get('Id'): r.get('Target') for r in rels}
    out = []
    for sh in wb.iter(f'{_SML_NS}sheet'):
        tgt = rid_target.get(sh.get(f'{_RNS}id'), '')
        if not tgt:
            continue
        tgt = tgt.lstrip('/') if tgt.startswith('/') else ('xl/' + tgt if not tgt.startswith('xl/') else tgt)
        out.append((sh.get('name'), tgt))
    return out


def _scan_workbook_xml(path: str):
    """One raw-XML pass over the workbook that returns (has_date_oor, error_map). It serves two safety
    needs calamine cannot: (1) DATE-OOR detection — a date-FORMATTED cell holding an out-of-range serial,
    which calamine mis-reads as a number where openpyxl errors → whole-file fallback; (2) ERROR-CELL
    recovery — calamine collapses every Excel error (#REF!/#VALUE!/…) to empty, losing its PRESENCE, which
    changes extraction (an error cell in a period column disambiguates plan/actual). The exact error string
    lives in the cell's <v>, so we harvest {sheet_name: {(ri,ci): '#REF!'}} to overlay back onto the grid,
    making error files byte-identical to openpyxl too. Raises on any structural problem so _parse_workbook
    fails closed to openpyxl. Early-exits to fallback the instant a date-OOR cell is seen."""
    err_map: dict = {}
    with zipfile.ZipFile(path) as z:
        date_styles = _date_format_style_ids(z)
        for name, xmlpath in _sheet_order(z):
            cells: dict = {}
            try:
                fh = z.open(xmlpath)
            except KeyError:
                err_map[name] = cells
                continue
            with fh:
                for _ev, el in _ET.iterparse(fh, events=('end',)):
                    if el.tag != f'{_SML_NS}c':
                        continue
                    t = el.get('t')
                    if t == 'e':
                        r = el.get('r')
                        if r:
                            cells[_coord_to_ri_ci(r)] = el.findtext(f'{_SML_NS}v') or ''
                    elif date_styles:
                        s = el.get('s')
                        if s is not None and int(s) in date_styles and t in (None, 'n'):
                            vt = el.findtext(f'{_SML_NS}v')
                            if vt:
                                try:
                                    x = float(vt)
                                except ValueError:
                                    x = None
                                if x is not None and (x > _MAX_XL_DATE_SERIAL or x < 0):
                                    return True, {}          # date-OOR → whole-file openpyxl fallback
                    el.clear()
            err_map[name] = cells
    return False, err_map


def _parse_workbook_calamine(path: str, err_map: dict):
    """calamine parse → normalized value grid (with Excel error cells overlaid back from err_map) +
    content/layout fingerprints built with the SAME formula as the openpyxl reference (absolute A1-origin
    coordinate == openpyxl's cell.coordinate; fingerprints iterate the FINAL overlaid grid in row-major
    order, exactly as openpyxl iterates), so clean AND error files are byte-identical at every level. Rows
    are padded to the sheet's width so the grid is rectangular like openpyxl's; an error cell beyond
    calamine's trimmed extent expands the grid so it is never lost. The OOR-date guard already cleared
    this file (else it fell back)."""
    with open(path, 'rb') as fh:
        byte_fp = hashlib.sha256(fh.read()).hexdigest()
    cwb = _CalamineWorkbook.from_path(path)
    content_parts: List[str] = []
    layout_parts: List[str] = []
    grid: dict = {}
    for sn in cwb.sheet_names:
        rows = [[_norm_calamine_value(v) for v in row]
                for row in cwb.get_sheet_by_name(sn).to_python(skip_empty_area=False)]
        for (ri, ci), err in err_map.get(sn, {}).items():   # restore error-cell PRESENCE (calamine drops it)
            while ri >= len(rows):
                rows.append([])
            if ci >= len(rows[ri]):
                rows[ri].extend([None] * (ci + 1 - len(rows[ri])))
            rows[ri][ci] = err
        width = max((len(r) for r in rows), default=0)
        grid[sn] = [r + [None] * (width - len(r)) for r in rows]
        content_parts.append(f'#SHEET#{sn}')
        layout_parts.append(f'#SHEET#{sn}')
        for ri, row in enumerate(grid[sn]):
            for ci, v in enumerate(row):
                if v is None or v == '':
                    continue
                coord = f'{_col_letter(ci + 1)}{ri + 1}'
                content_parts.append(f'{coord}={v!r}')
                if not _is_volatile_value(v):
                    layout_parts.append(f'{coord}={str(v).strip().lower()!r}')
    content_fp = _sha(content_parts) if content_parts else byte_fp
    layout_fp = _sha(layout_parts) if layout_parts else byte_fp
    return (byte_fp, content_fp, layout_fp, grid, None)


def _parse_workbook(path: str):
    """Single workbook parse feeding BOTH identity and the profiler grid. Primary reader = calamine
    (fast); a whole-file openpyxl fallback fires when calamine is unavailable, when the date-OOR guard
    trips, or on ANY calamine/scan error (fail-closed). The scan also harvests Excel error cells so the
    calamine grid can restore their presence (byte-identical to openpyxl). openpyxl stays the byte-exact
    reference for fallen-back files. Returns (byte_fp, content_fp, layout_fp, grid, error) like the ref."""
    if not _HAS_CALAMINE:
        return _parse_workbook_openpyxl(path)
    try:
        oor, err_map = _scan_workbook_xml(path)
        if oor:
            _FALLBACK_LOG.append((path, 'date_oor'))
            return _parse_workbook_openpyxl(path)
        return _parse_workbook_calamine(path, err_map)
    except Exception as e:
        _FALLBACK_LOG.append((path, f'calamine_error:{type(e).__name__}'))
        return _parse_workbook_openpyxl(path)


def _parse_workbook_cached(path: str):
    """Per-run memo of _parse_workbook, keyed by (path, mtime, size) like the profiler — one parse
    per file per run, shared by identity + profiling. Cleared each run by clear_parse_cache."""
    try:
        st = os.stat(path)
        key = (path, st.st_mtime_ns, st.st_size)
    except OSError:
        key = (path, None, None)
    hit = _PARSE_CACHE.get(key)
    if hit is not None:
        return hit
    r = _parse_workbook(path)
    _PARSE_CACHE[key] = r
    return r


def compute_identity(label: str, path: str) -> FileIdentity:
    """Compute all three fingerprints from the single unified parse (data_only so cached computed
    values are read, never formulas — ALWAYS data_only=True, NEVER eval a formula). The same parse
    also yields the profiler grid, cached per-run so profile_file reuses this exact open (parse-once)."""
    with open(path, 'rb') as fh:
        byte_fp = hashlib.sha256(fh.read()).hexdigest()

    cached = _IDENTITY_CACHE.get(byte_fp)          # pure, byte-keyed → any edit misses, never stale
    if cached is not None:
        _IDENTITY_CACHE.move_to_end(byte_fp)       # most-recently-used (LRU bookkeeping)
        return FileIdentity(label, path, byte_fp, cached[0], cached[1])

    _bfp, content_fp, layout_fp, _grid, err = _parse_workbook_cached(path)
    _cache_put(byte_fp, content_fp, layout_fp)
    if err is not None:                            # non-workbook/unreadable → identity by bytes
        return FileIdentity(label, path, byte_fp, byte_fp, byte_fp, ())
    return FileIdentity(label=label, path=path, byte_fp=byte_fp,
                        content_fp=content_fp, layout_fp=layout_fp)


def canonical_sort(identities: List[FileIdentity]) -> List[FileIdentity]:
    """Deterministic processing order — by content fingerprint, tie-broken by
    label. Upload order and filenames never influence the output."""
    return sorted(identities, key=lambda i: (i.content_fp, i.label))


def _hash_py_dir(dirpath: str) -> str:
    """Hash every *.py directly in `dirpath` (sorted, top-level only — no recursion into a
    tests/ subdir or __pycache__). Pure + deterministic so it can be negative-controlled:
    editing OR adding a module changes the digest. Under-coverage is impossible by
    construction — it hashes the whole directory, never a hand-list."""
    h = hashlib.sha256()
    try:
        names = sorted(f for f in os.listdir(dirpath) if f.endswith('.py'))
    except OSError:
        names = []
    for name in names:
        h.update(name.encode())                    # include the name → adding a module invalidates
        try:
            with open(os.path.join(dirpath, name), 'rb') as fh:
                h.update(fh.read())
        except OSError:
            h.update(b'?missing?')
        h.update(b'\x1f')
    return 'lv_' + h.hexdigest()[:16]


def net_logic_version() -> str:
    """Structural code-version for the extraction cache key: a hash of EVERY library module's
    source (see _hash_py_dir). ANY change to library logic yields a new version → cached results
    computed under the old code are missed and recomputed, so a shipped fix can never be masked
    by a stale cache hit. Whole-package by design: a hand-picked 'value modules' list silently
    missed 6 modules when traced (see module note) — over-inclusion is a safe false-miss,
    under-inclusion is a stale-logic hole. Computed once per process (source is immutable)."""
    global _LOGIC_VERSION
    if _LOGIC_VERSION is None:
        _LOGIC_VERSION = _hash_py_dir(os.path.dirname(os.path.abspath(__file__)))
    return _LOGIC_VERSION


def extraction_cache_key(content_fp: str, *, as_of, rate_card_id, anchor_cr,
                         domicile, use_model, base_currency=None, logic_version: str = None,
                         contract_sig: str = None) -> str:
    """The PROVABLY-COMPLETE key for the per-file extraction/reuse cache: content_fp PLUS
    every run input that can change the extracted CIR. content_fp ALONE is a config-staleness
    bug — a file's INR-config record served under a new rate card / as_of / anchor / domicile
    (goes live at the U6 multi-currency boundary). `entity` is EXCLUDED: empirically proven
    pure attribution (0 value-diffs when it alone changes), so keying on it would only cause
    false cache misses. `as_of` is kept explicit even though today it reaches extraction only
    via rate_card_id — future-proofing as_of-dependent period selection. `base_currency` (the
    fund's user-confirmed base) is a THIRD positive-evidence source at resolve_currency: it can
    flip a no-evidence HOLD into a base-currency EMIT, so it changes the CIR and MUST be in the
    key — else a file extracted (held) under base=None would be served STALE on a re-run once the
    fund base is supplied (the exact Inc-5 stale-config class, now for the currency dimension).
    logic_version defaults to the structural net-module hash; a caller may inject one (tests / a
    pinned prod version). contract_sig (schema + SEED-and-LEARNED lexicon + identities + checks +
    tolerances) closes D8: net_logic_version hashes preingest3 .py source only, so it misses the
    out-of-package OUTPUT SCHEMA and the runtime-grown learned-synonym JSON — both change the CIR."""
    lv = logic_version if logic_version is not None else net_logic_version()
    if contract_sig is None:
        from .contract import contract_signature      # deferred: identity is foundational
        contract_sig = contract_signature()
    return _sha([content_fp, str(as_of), str(rate_card_id),
                 str(anchor_cr), str(domicile), '1' if use_model else '0',
                 str(base_currency), lv, contract_sig])
