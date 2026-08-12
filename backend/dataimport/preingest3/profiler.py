"""
S2 — Profiling (the "chunking" mechanism, done the v2 way).

Revision 1.0 chunked raw rows into many model calls; this design REPLACES that.
The model never sees the workbook — only a compact structural map per sheet:
sheet names, dimensions, detected header rows, the label column, detected table
blocks, and a BOUNDED WINDOW of representative cells WITH THEIR ADDRESSES
(Section 05). "A fourteen-megabyte file with a junk sheet of sixteen thousand
columns reduces to a few kilobytes … noise is removed before a token is spent."

Two rules from the design govern the reduction:
  • "Sample large regions, never load them" (S2) — the window is bounded.
  • "Ledgers are aggregated in code … a six-thousand-row transaction listing is
    summed by code and never sent" (Section 17) — a raw numeric ledger is flagged
    is_raw_ledger=True so the extractor code-aggregates it instead of locating in
    it. This is what makes row-explosion structurally impossible.

Everything here is deterministic code — zero tokens. The profile it emits is the
sole input the locator (S4) ever sees.
"""
from __future__ import annotations

import datetime as _dt
import functools
import os
import re
from dataclasses import dataclass, field
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.utils import get_column_letter


def token_present(tok: str, low: str) -> bool:
    """A hint token appears as a WHOLE token in `low` (already lower-cased), never
    as a substring — the single guard against the short-token substring bug class:
    unit 'cr' must match 'in ₹ cr' / '(cr)' but NOT 'des-cr-iption' or 'a-cr-oss';
    currency 'rs' must not match 'hou-rs'. Alnum boundaries only, so currency
    symbols (₹ $ €) and the "'000" unit still match. ONE implementation, reused by
    the profiler's hint scan and the extractor's local unit/currency read."""
    if not tok:
        return False
    if not tok[0].isalnum() and not tok[-1].isalnum():   # pure symbol like ₹ $ €
        return tok in low
    return re.search(r'(?<![a-z0-9])' + re.escape(tok) + r'(?![a-z0-9])', low) is not None

# Bounds — the window is capped so a giant sheet still profiles to a few KB.
MAX_LABELS = 150        # label-column entries shown to the model
MAX_HEADERS = 40        # header cells shown
MAX_SAMPLE = 60         # representative data cells sampled
HEADER_SCAN_ROWS = 20   # detect the header within the first N rows
LEDGER_ROW_THRESHOLD = 60      # a sheet longer than this is a record-listing candidate
LEDGER_LABEL_CARDINALITY = 0.6 # label-column distinct/total ratio marking homogeneous records

# Currency / unit hints scanned from header & title text (feeds Rate Card
# coverage and Quantity.scale). Universal words, never fund-specific.
_CCY_HINTS = {
    '₹': 'INR', 'rs': 'INR', 'inr': 'INR', 'rupee': 'INR',
    'rm': 'MYR', 'myr': 'MYR', 'ringgit': 'MYR',
    '$': 'USD', 'usd': 'USD', 'dollar': 'USD',
    '€': 'EUR', 'eur': 'EUR', '£': 'GBP', 'gbp': 'GBP',
    'sgd': 'SGD', 'aed': 'AED',
}
# singular AND plural forms: token_present matches whole tokens, so 'million' does
# NOT match 'millions' (the trailing 's' breaks the boundary) — a declared
# '(₹ Millions)' header would otherwise be missed and the whole statement held.
_UNIT_HINTS = ['crores', 'crore', 'cr', 'lakhs', 'lakh', 'lacs', 'lac', 'millions', 'million',
               'mn', 'thousands', 'thousand', "'000", 'billions', 'billion', 'bn']


def _cell_type(v) -> str:
    if v is None or v == '':
        return 'blank'
    if isinstance(v, bool):
        return 'bool'
    if isinstance(v, (int, float)):
        return 'num'
    if isinstance(v, (_dt.datetime, _dt.date, _dt.time)):
        return 'date'
    return 'text'


def _addr(col_idx0: int, row_idx0: int) -> str:
    """0-based indices → A1 address."""
    return f'{get_column_letter(col_idx0 + 1)}{row_idx0 + 1}'


@dataclass
class SheetProfile:
    sheet: str
    n_rows: int
    n_cols: int
    header_rows: List[int] = field(default_factory=list)   # 0-based
    label_col: Optional[int] = None                        # 0-based
    table_blocks: List[dict] = field(default_factory=list)
    labels: List[dict] = field(default_factory=list)       # [{addr,text}]
    headers: List[dict] = field(default_factory=list)      # [{addr,text}]
    sample: List[dict] = field(default_factory=list)       # [{addr,value,type}]
    currency_hints: List[str] = field(default_factory=list)
    unit_hints: List[str] = field(default_factory=list)
    is_raw_ledger: bool = False
    token_estimate: int = 0

    def to_model_view(self) -> dict:
        """The compact, addressed structure the locator receives — never values
        in bulk, only the navigational axes + a bounded sample."""
        return {
            'sheet': self.sheet, 'dims': [self.n_rows, self.n_cols],
            'header_rows': self.header_rows, 'label_col': self.label_col,
            'table_blocks': self.table_blocks,
            'labels': self.labels[:MAX_LABELS], 'headers': self.headers[:MAX_HEADERS],
            'sample': self.sample[:MAX_SAMPLE],
            'currency_hints': self.currency_hints, 'unit_hints': self.unit_hints,
            'is_raw_ledger': self.is_raw_ledger,
        }


def _read_grid(path: str) -> Dict[str, List[List[Any]]]:
    """Load every sheet as a 2-D list of values (data_only — cached values,
    never formulas)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    grid: Dict[str, List[List[Any]]] = {}
    for sn in wb.sheetnames:
        ws = wb[sn]
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        grid[sn] = rows
    try:
        wb.close()
    except Exception:
        pass
    return grid


def _detect_header_row(rows: List[List[Any]]) -> Optional[int]:
    """The row (within the first HEADER_SCAN_ROWS) with the most text cells and
    at least two — a universal structural heuristic, never a hardcoded row."""
    best, best_score = None, 0
    for r in range(min(HEADER_SCAN_ROWS, len(rows))):
        texts = sum(1 for v in rows[r] if _cell_type(v) == 'text')
        if texts >= 2 and texts > best_score:
            best, best_score = r, texts
    return best


def _detect_label_col(rows: List[List[Any]], header_row: Optional[int]) -> Optional[int]:
    """The column with the most text cells below the header — the entity/line-
    item axis a locator walks left to."""
    start = (header_row + 1) if header_row is not None else 0
    n_cols = max((len(r) for r in rows), default=0)
    best, best_score = None, 0
    for c in range(min(n_cols, 8)):   # label column is near the left in practice
        texts = sum(1 for r in range(start, len(rows))
                    if c < len(rows[r]) and _cell_type(rows[r][c]) == 'text')
        if texts > best_score:
            best, best_score = c, texts
    return best


def _detect_blocks(rows: List[List[Any]], header_row, label_col) -> List[dict]:
    """Segment a sheet into contiguous data blocks. Blocks are separated by
    fully-blank rows, but reading NEVER stops on a blank (critical rule) — a
    block simply ends and the next begins. Each block records its own row span."""
    blocks, start = [], None
    for r in range(len(rows)):
        blank = all(_cell_type(v) == 'blank' for v in rows[r]) if rows[r] else True
        if not blank and start is None:
            start = r
        elif blank and start is not None:
            blocks.append({'start_row': start, 'end_row': r - 1})
            start = None
    if start is not None:
        blocks.append({'start_row': start, 'end_row': len(rows) - 1})
    return blocks


def _dominant_col_type(rows: List[List[Any]], col: int, start: int) -> str:
    """Majority cell type in a column below the header — how the sheet uses that
    column, independent of any specific file's headers."""
    counts = {'num': 0, 'text': 0, 'date': 0}
    for r in range(start, len(rows)):
        if col < len(rows[r]):
            t = _cell_type(rows[r][col])
            if t in counts:
                counts[t] += 1
    return max(counts, key=counts.get) if any(counts.values()) else 'blank'


def _is_raw_ledger(rows: List[List[Any]], header_row, label_col) -> bool:
    """Universal discriminator between a 'record listing' (code-aggregate, never
    send) and a 'named-metric statement' (keep, locate directly).

    A ledger is defined STRUCTURALLY, not by numeric density: it is a LONG list
    of HOMOGENEOUS records — the same column schema repeated over many rows,
    whose label column is high-cardinality (unique txn ids / dates / names)
    rather than a bounded vocabulary of business concepts. A P&L has ~15–40
    rows of DISTINCT NAMED concepts; a general ledger has thousands of near-
    unique instances. This holds for a 2-column ledger and a 12-column one
    alike, and for files not yet seen — no threshold tuned to one sheet.
    """
    start = (header_row + 1) if header_row is not None else 0
    data_rows = [r for r in range(start, len(rows))
                 if any(_cell_type(v) != 'blank' for v in rows[r])]
    if len(data_rows) < LEDGER_ROW_THRESHOLD:
        return False   # short enough to locate directly

    n_cols = max((len(rows[r]) for r in data_rows), default=0)
    numeric_cols = sum(1 for c in range(n_cols)
                       if _dominant_col_type(rows, c, start) == 'num')
    if numeric_cols == 0:
        return False   # nothing to aggregate — treat as a (long) text list, still locate

    # Cardinality of the label axis: unique labels / populated labels.
    if label_col is None:
        return True    # long, numeric, no clear metric axis ⇒ record grid
    labels = [str(rows[r][label_col]).strip().lower()
              for r in data_rows
              if label_col < len(rows[r]) and _cell_type(rows[r][label_col]) == 'text']
    if not labels:
        return True    # long numeric grid with no textual metric names ⇒ ledger
    distinct_ratio = len(set(labels)) / len(labels)
    return distinct_ratio >= LEDGER_LABEL_CARDINALITY


def _scan_hints(rows: List[List[Any]]):
    ccy, unit = set(), set()
    for r in range(min(HEADER_SCAN_ROWS, len(rows))):
        for v in rows[r]:
            if _cell_type(v) != 'text':
                continue
            low = str(v).lower()
            for token, code in _CCY_HINTS.items():
                if token_present(token, low):
                    ccy.add(code)
            for u in _UNIT_HINTS:
                if token_present(u, low):
                    unit.add(u)
    return sorted(ccy), sorted(unit)


def _build_windows(rows, header_row, label_col):
    labels, headers, sample = [], [], []
    if header_row is not None:
        for c, v in enumerate(rows[header_row]):
            if _cell_type(v) == 'text':
                headers.append({'addr': _addr(c, header_row), 'text': str(v).strip()})
    start = (header_row + 1) if header_row is not None else 0
    if label_col is not None:
        for r in range(start, len(rows)):
            if label_col < len(rows[r]):
                v = rows[r][label_col]
                if _cell_type(v) == 'text':
                    labels.append({'addr': _addr(label_col, r), 'text': str(v).strip()})
    # representative data sample — spread across the sheet, bounded
    data_cells = []
    for r in range(start, len(rows)):
        for c, v in enumerate(rows[r]):
            if c == label_col:
                continue
            t = _cell_type(v)
            if t in ('num', 'date'):
                data_cells.append({'addr': _addr(c, r),
                                   'value': v.isoformat() if t == 'date' else v,
                                   'type': t})
    if len(data_cells) <= MAX_SAMPLE:
        sample = data_cells
    else:
        step = len(data_cells) / MAX_SAMPLE
        sample = [data_cells[int(i * step)] for i in range(MAX_SAMPLE)]
    return labels, headers, sample


def profile_sheet(name: str, rows: List[List[Any]]) -> SheetProfile:
    n_rows = len(rows)
    n_cols = max((len(r) for r in rows), default=0)
    header_row = _detect_header_row(rows)
    label_col = _detect_label_col(rows, header_row)
    blocks = _detect_blocks(rows, header_row, label_col)
    raw_ledger = _is_raw_ledger(rows, header_row, label_col)
    ccy, unit = _scan_hints(rows)
    labels, headers, sample = ([], [], []) if raw_ledger else _build_windows(rows, header_row, label_col)
    p = SheetProfile(
        sheet=name, n_rows=n_rows, n_cols=n_cols,
        header_rows=[header_row] if header_row is not None else [],
        label_col=label_col, table_blocks=blocks,
        labels=labels, headers=headers, sample=sample,
        currency_hints=ccy, unit_hints=unit, is_raw_ledger=raw_ledger,
    )
    import json
    p.token_estimate = len(json.dumps(p.to_model_view(), default=str)) // 4
    return p


def _profile_core_uncached(path: str) -> dict:
    """The label-independent, deterministic parse of one workbook: the expensive work
    (open + materialise every sheet into plain value-lists + per-sheet profiling). Kept
    separate from `label` so it can be memoised and shared across the router, anchor
    builder, per-company extractor and fund-terms reader, which each ask for the same
    file. Returns plain-list grids (never openpyxl objects), so every consumer only reads."""
    try:
        grid = _read_grid(path)
    except Exception as e:  # noqa: BLE001 — a corrupt/empty/non-workbook file must
        # not crash a 15-file run; it becomes a disclosed, sheet-less profile that
        # every downstream reader already treats as "nothing to bind" (fail-closed).
        return {'path': path, 'sheets': [], 'grid': {},
                'currencies': [], 'error': f'{type(e).__name__}: {e}'[:120]}
    sheets = [profile_sheet(sn, rows) for sn, rows in grid.items()]
    currencies = sorted({c for s in sheets for c in s.currency_hints})
    return {'path': path, 'sheets': sheets, 'grid': grid,
            'currencies': currencies, 'error': None}


@functools.lru_cache(maxsize=None)
def _profile_core_cached(path: str, _mtime_ns: int, _size: int) -> dict:
    """Memoised by (path, mtime, size): one parse per file per run, re-parsed only if the
    file's bytes change. clear_profile_cache() resets it at each pipeline run so memory
    never accumulates across runs. The result is READ-ONLY shared — never mutate it."""
    return _profile_core_uncached(path)


def clear_profile_cache() -> None:
    """Drop the parse cache. Called once at the start of a pipeline run so a run reuses each
    file's grid but no run ever holds another run's grids. Correctness-neutral (forces a
    re-parse to an identical grid), memory-bounding only."""
    _profile_core_cached.cache_clear()


def profile_file(label: str, path: str) -> dict:
    """Return {'label','path','sheets':[SheetProfile...],'grid':{sheet:rows},
    'currencies':[...],'error'} — the grid is kept in-process for the code-side reader
    (S5) to open located cells; only the bounded model_view ever reaches S4.

    The parse is memoised by (path, mtime, size) so the same file is opened ONCE per run
    and reused by every stage; `label` is applied per-call as a fresh top-level dict over
    the shared read-only core, so each caller keeps its own label with no cross-talk."""
    try:
        st = os.stat(path)
        core = _profile_core_cached(path, st.st_mtime_ns, st.st_size)
    except OSError:  # path missing/unstatable → reproduce the exact uncached error-profile
        core = _profile_core_uncached(path)
    return {'label': label, **core}
