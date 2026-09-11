"""Two root-cause fixes for the GUI report (workbook shape + currency ping-pong), each with a
negative control that reddens on the exact bug.

FIX 1 — the Download delivers the CONSOLIDATED MASTER workbook, not the thin review-gate one.
  The frontend saved `assemble.build()` (10 review sheets: Cover/Portfolio_KPI/_Provenance/…), so the
  downloaded file was missing NAV_CALC, CAPITAL_CALLS, WATERFALL, DASHBOARD_BRIDGE, etc. — even though
  the engine produced them. _run_job now builds `master_workbook.build_master()` (the full canonical
  15-tab product). Reddening: against the old wiring the master-only sheets are absent → the assertion
  goes RED. Driven through the REAL endpoints + REAL worker (production path, deterministic/model-off).

FIX 2 — the Rate Card ACCUMULATES across the 'hold → supply rate → re-run' loop.
  /ratecard/ overwrote the stored card, so supplying the 2nd currency dropped the 1st and the run
  re-reported it uncovered — an endless SGD↔MYR prompt that could never cover both. It now merges.
  Reddening: against the overwrite the stored card keeps only the last currency → the 'both present'
  assertion goes RED. Endpoint-level, no pipeline, token-free.
"""
import io
import os

import openpyxl
from django.conf import settings
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from rest_framework.test import APIClient

from accounts.models import Organization, User
from dataimport import preingest3_views
from dataimport.models import PreIngestJob

_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

_FUND = [['Company', 'Cost', 'Fair Value', 'Ownership %', 'Country'],
         ['Zephyr Diagnostics', 60, 150, 18, 'India']]
_MIS = [['Particulars', 'Apr-25', 'May-25', 'Jun-25'],
        ['Revenue', 3000000, 3100000, 3200000],
        ['EBITDA', 500000, 510000, 520000],
        ['Closing Cash', 4800000, 4900000, 5000000]]

# the full consolidated product; a few master-only tabs the thin review-gate workbook never had.
_MASTER_ONLY = {'MASTER_INPUTS', 'CAPITAL_CALLS', 'NAV_CALC', 'MOIC_TVPI_DPI',
                'WATERFALL_EUR', 'DASHBOARD_BRIDGE', 'PORTFOLIO_MASTER'}
_REVIEW_GATE_ONLY = {'Cover', '_Provenance', '_Disclosures', '_Audit'}


def _xlsx(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class DownloadShapeTest(TestCase):
    def setUp(self):
        self.org = Organization.objects.create(slug='dlshape', name='DL Shape Org')
        self.user = User.objects.create(username='gp-dl', role='gp_admin', organization=self.org)
        self.client = APIClient()
        self.client.force_authenticate(self.user)
        self._orig_kick = preingest3_views._kick
        preingest3_views._kick = lambda job_id: preingest3_views._run_job(str(job_id))

    def tearDown(self):
        preingest3_views._kick = self._orig_kick

    def test_download_is_the_consolidated_master_workbook(self):
        fund = SimpleUploadedFile('fund.xlsx', _xlsx(_FUND), content_type=_XLSX)
        mis = SimpleUploadedFile('Zephyr.xlsx', _xlsx(_MIS), content_type=_XLSX)
        r = self.client.post(reverse('preingest3-upload'), {'files': [fund, mis]}, format='multipart')
        self.assertEqual(r.status_code, 201, r.content)
        job = PreIngestJob.objects.get(pk=r.data['job_id'])
        self.assertIn(job.status, ('completed', 'completed_with_errors'), job.progress_message)
        self.assertTrue(job.output_file, 'a workbook must be produced')

        path = os.path.join(settings.MEDIA_ROOT, job.output_file)
        wb = openpyxl.load_workbook(path, read_only=True)
        sheets = set(wb.sheetnames)
        wb.close()
        # the master product tabs are present (RED against the old assemble.build wiring)…
        self.assertTrue(_MASTER_ONLY.issubset(sheets),
                        f'missing master tabs: {sorted(_MASTER_ONLY - sheets)} — got {sorted(sheets)}')
        # …and it is NOT the thin review-gate workbook.
        self.assertFalse(_REVIEW_GATE_ONLY & sheets,
                         f'download is still the review-gate workbook: {sorted(_REVIEW_GATE_ONLY & sheets)}')


class RateCardMergeTest(TestCase):
    def setUp(self):
        self.org = Organization.objects.create(slug='rcmerge', name='RC Merge Org')
        self.user = User.objects.create(username='gp-rc', role='gp_admin', organization=self.org)
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def _job_needing(self, currencies):
        # a job whose last run reported these currencies uncovered (drives the intake actionability).
        return PreIngestJob.objects.create(
            organization=self.org, uploaded_by=self.user, total_files=1, status='completed',
            summary={'engine': preingest3_views._ENGINE, 'as_of': '2026-09-10',
                     'currency_report': {'uncovered': [{'currency': c} for c in currencies]}})

    def _supply(self, job, currency, rate):
        return self.client.post(
            reverse('preingest3-ratecard', args=[str(job.id)]),
            {'manual_rates': [{'currency': currency, 'rate': rate,
                               'date': '2026-09-10', 'source': 'test'}]}, format='json')

    def test_two_currencies_supplied_one_at_a_time_both_persist(self):
        job = self._job_needing(['SGD', 'MYR'])
        r1 = self._supply(job, 'SGD', '75.36')
        self.assertEqual(r1.status_code, 200, r1.content)
        r2 = self._supply(job, 'MYR', '23.47')
        self.assertEqual(r2.status_code, 200, r2.content)

        job.refresh_from_db()
        stored = {row['currency'] for row in (job.summary['rate_card']['rates'])}
        # BOTH must survive (RED against the old overwrite, which kept only MYR)…
        self.assertEqual(stored, {'SGD', 'MYR'}, f'rate card did not accumulate — stored {stored}')
        # …and the 2nd response reports nothing still uncovered (the ping-pong is broken).
        self.assertEqual(r2.data['still_uncovered'], [])

    def test_resupplying_a_currency_updates_its_rate_not_drops_the_other(self):
        job = self._job_needing(['SGD', 'MYR'])
        self._supply(job, 'SGD', '75.36')
        self._supply(job, 'MYR', '23.47')
        self._supply(job, 'SGD', '76.10')            # correct a rate
        job.refresh_from_db()
        rows = {row['currency']: row['inr_per_unit'] for row in job.summary['rate_card']['rates']}
        self.assertEqual(set(rows), {'SGD', 'MYR'})   # MYR still there
        self.assertEqual(str(rows['SGD']), '76.10')   # SGD updated to the new value
